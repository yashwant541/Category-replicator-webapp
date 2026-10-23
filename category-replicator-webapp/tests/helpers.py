"""Synthetic example generator shared by the tests and the dev server."""
import datetime as dt
import io
import math
import random


def xl_round(x, d=0):
    f = 10 ** d
    return math.floor(abs(x) * f + 0.5 + 1e-9) / f * (1 if x >= 0 else -1)

import openpyxl


def make_case(n=60, seed=1, with_formulas=False, split=True):
    rnd = random.Random(seed)
    regions = ["EU", "US", "ASIA"]
    tiers = {"AAA": 0.02, "AA": 0.035, "A": 0.05, "BBB": 0.08}
    country_group = {"FR": "Core", "DE": "Core", "US": "Americas", "JP": "APAC", "IN": "APAC"}
    rows = []
    for i in range(n):
        country = rnd.choice(list(country_group))
        region = {"FR": "EU", "DE": "EU", "US": "US", "JP": "ASIA", "IN": "ASIA"}[country]
        rows.append(dict(
            ID=f"T{1000 + i}", Region=region, Country=country, Tier=rnd.choice(list(tiers)),
            Amount=rnd.choice([250, 800, 1200, 5400, 9999.5, 15000]) + rnd.randint(0, 99),
            Crm=rnd.choice([0, 0.1, 0.25]), Date=dt.datetime(rnd.choice([2023, 2024, 2025]), rnd.randint(1, 12), rnd.randint(1, 28)),
        ))
    # --- input workbook
    inp = openpyxl.Workbook()
    ws = inp.active
    ws.title = "Data"
    ws.append(["ID", "Region", "Country", "Tier", "Amount", "Crm", "Date"])
    for r in rows:
        ws.append([r[k] for k in ["ID", "Region", "Country", "Tier", "Amount", "Crm", "Date"]])
    # --- mapping workbook
    mp = openpyxl.Workbook()
    m1 = mp.active
    m1.title = "Rates"
    m1.append(["Tier", "Rate"])
    for t, v in tiers.items():
        m1.append([t, v])
    m2 = mp.create_sheet("Groups")
    m2.append(["Country", "Group"])
    for c, g in country_group.items():
        m2.append([c, g])

    def derive(r):
        rate = tiers[r["Tier"]]
        fee = xl_round(r["Amount"] * rate, 2)
        return dict(ID=r["ID"], Region=r["Region"], Amount=r["Amount"], Group=country_group[r["Country"]],
                    Rate=rate, Fee=fee, Net=xl_round(r["Amount"] - fee, 2),
                    RWA=xl_round(r["Amount"] * (1 - r["Crm"]), 2),
                    Bucket="High" if r["Amount"] >= 1000 else "Low",
                    Label=f"{r['Region']}-{r['ID']}", Year=r["Date"].year,
                    Name=r["Tier"].lower())
    out_rows = [derive(r) for r in rows]
    cols = list(out_rows[0])
    out = openpyxl.Workbook()
    out.remove(out.active)
    if split:
        for reg in regions:
            w = out.create_sheet(reg)
            w.append(cols)
            for o in out_rows:
                if o["Region"] == reg:
                    w.append([o[c] for c in cols])
    else:
        w = out.create_sheet("Result")
        w.append(cols)
        for o in out_rows:
            w.append([o[c] for c in cols])
    def b(wb):
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()
    return b(inp), b(mp), b(out), rows, out_rows
