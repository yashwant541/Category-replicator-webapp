import io
import tempfile

import openpyxl
from replicator import service
from replicator.store import LocalStore


def wb_bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_files():
    # raw input: one sheet per segment, NO derived columns, no split column
    inp = openpyxl.Workbook()
    inp.remove(inp.active)
    data = {"EU": [("E1", "A", 100), ("E2", "B", 200), ("E3", "A", 300)],
            "US": [("U1", "B", 1000), ("U2", "A", 500)]}
    for sh, rows in data.items():
        w = inp.create_sheet(sh)
        w.append(["ID", "Tier", "Amount"])
        for r in rows:
            w.append(list(r))
    # mapping workbook: several sheets
    mp = openpyxl.Workbook()
    r = mp.active
    r.title = "Rates"
    r.append(["Tier", "Rate"]); r.append(["A", 0.1]); r.append(["B", 0.2])
    g = mp.create_sheet("Caps")
    g.append(["Tier", "Cap"]); g.append(["A", 50]); g.append(["B", 150])
    # output workbook: data sheets WITH formulas (EU and US use different Fee formulas) + the mapping sheets
    out = openpyxl.Workbook()
    out.remove(out.active)
    for sh, rows in data.items():
        w = out.create_sheet(sh)
        w.append(["ID", "Tier", "Amount", "Rate", "Fee", "Capped"])
        for i, (a, b, c) in enumerate(rows, start=2):
            fee = f"=C{i}*D{i}" if sh == "EU" else f"=ROUND(C{i}*D{i}*1.1,1)"
            w.append([a, b, c, f"=VLOOKUP(B{i},Rates!$A$2:$B$3,2,FALSE)", fee,
                      f"=MIN(E{i},VLOOKUP(B{i},Caps!$A$2:$B$3,2,FALSE))"])
    r2 = out.create_sheet("Rates")
    r2.append(["Tier", "Rate"]); r2.append(["A", 0.1]); r2.append(["B", 0.2])
    g2 = out.create_sheet("Caps")
    g2.append(["Tier", "Cap"]); g2.append(["A", 50]); g2.append(["B", 150])
    return wb_bytes(inp), wb_bytes(mp), wb_bytes(out)


def test_split_input_sheets_formulas_per_sheet():
    inp, mp, out = make_files()
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    service.add_draft_files(st, did, [("in.xlsx", inp)], "input")
    service.add_draft_files(st, did, [("map.xlsx", mp)], "mapping")
    d = service.add_draft_files(st, did, [("out.xlsx", out)], "output")
    roles = {s["sheet"]: s["role"] for s in d["files"][2]["sheets"]}
    assert roles == {"EU": "output", "US": "output", "Rates": "mapping", "Caps": "mapping"}, roles
    s = service.create_system(st, did, {"name": "seg"})
    assert s["input"]["columns"][-1] == "Source sheet"
    sp = s["filters"]["split"]
    assert sp["enabled"] and sp["column"] == "Source sheet"
    by = {r["name"]: r for r in s["rules"]}
    assert by["Fee"]["by_sheet"].keys() == {"EU", "US"}
    assert by["Rate"]["kind"] == "formula" and not by["Rate"].get("by_sheet")
    assert set(s["tables"]) == {"Rates", "Caps"}          # no duplicates from the output workbook

    # NEW data: raw input only, again one sheet per segment
    new = openpyxl.Workbook()
    new.remove(new.active)
    for sh, rows in {"EU": [("E9", "B", 400)], "US": [("U9", "A", 700), ("U8", "B", 10)]}.items():
        w = new.create_sheet(sh)
        w.append(["ID", "Tier", "Amount"])
        for r in rows:
            w.append(list(r))
    run = service.run_system(st, s["id"], [("new.xlsx", wb_bytes(new))])
    assert set(run["sheets"]) == {"EU", "US"}
    eu = dict(zip(run["preview"]["EU"]["headers"], run["preview"]["EU"]["data"][0]))
    assert eu["Rate"] == 0.2 and abs(eu["Fee"] - 80) < 1e-9 and eu["Capped"] == 80
    us = [dict(zip(run["preview"]["US"]["headers"], r)) for r in run["preview"]["US"]["data"]]
    assert abs(us[0]["Fee"] - 77) < 1e-9 and us[0]["Capped"] == 50   # US formula ROUND(C*D*1.1,1), cap A = 50
    assert us[1]["Fee"] == 2.2 and us[1]["Capped"] == 2.2
    rp = service.replay(st, s["id"])
    assert not rp["rule_errors"], rp["rule_errors"]
