"""Locking the data type of chosen output columns in a system."""
import datetime as dt

import openpyxl
import pytest

from replicator import engine, service
from test_template import b, make, output_book

# sample: ID is text-looking, Amt is a decimal, Day a date, Qty whole numbers, Fee a formula
INPUT = [("A001", 12.5, dt.datetime(2026, 1, 5), 3), ("A002", 99.999, dt.datetime(2026, 2, 6), 4)]


def book(rows, header=("ID", "Amt", "Day", "Qty")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    return b(wb)


def system():
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Amt", "Day", "Qty", "Fee"])
    for i, r in enumerate(INPUT, start=2):
        o.append(list(r) + [f"=B{i}*0.1"])
    return make([("i.xlsx", book(INPUT), "input"), ("o.xlsx", b(out), "output")])


def col(s, name):
    return next(c for c in s["sheets"][0]["columns"] if c["name"] == name)


def lock(st, s, *edits):
    return service.update_system(st, s["id"], {"dtypes": [dict(sheet="Out", **e) for e in edits]})


def test_the_samples_own_types_are_detected():
    st, s = system()
    assert [col(s, c)["sample_type"] for c in ("ID", "Amt", "Day", "Qty")] == ["text", "decimal", "date", "integer"]


def test_coerce_value_unit():
    cv = engine.coerce_value
    assert cv(123, {"type": "text"}) == ("123", True) and cv(7.0, {"type": "text"}) == ("7", True)
    assert cv("1,234.6", {"type": "integer"}) == (1235, True) and cv(-2.5, {"type": "integer"}) == (-3, True)
    assert cv("12.345", {"type": "decimal", "decimals": 2}) == (12.35, True)
    assert cv("abc", {"type": "integer"}) == (None, False) and cv(None, {"type": "integer"}) == (None, True)
    assert cv("2026-03-04", {"type": "date"})[0] == dt.datetime(2026, 3, 4)
    assert cv(46000, {"type": "date"})[0] == dt.datetime(2025, 12, 9)            # Excel serial number
    assert cv("not a date", {"type": "date"}) == (None, False) and cv(5, {"type": "date"}) == (None, False)


def test_locked_types_are_enforced_on_new_input_and_formatted_in_excel():
    st, s = system()
    s = lock(st, s, dict(column="ID", type="text"), dict(column="Amt", type="decimal", decimals=1),
             dict(column="Day", type="date", format="dd-mmm-yyyy"), dict(column="Qty", type="integer"),
             dict(column="Fee", type="decimal", decimals=3))
    assert col(service.get_system(st, s["id"]), "Day")["dtype"] == {"type": "date", "format": "dd-mmm-yyyy"}   # saved
    # new input arrives "dirty": the ID as a number, amounts as text with commas, dates as text, quantity as decimals
    new = book([(77, "1,234.56", "2026-03-04", "5.4"), (88, "9", "04/05/2026", "oops")])
    run, wb = output_book(st, s, new)
    ws = wb["Out"]
    assert [ws["A2"].value, ws["B2"].value, ws["C2"].value, ws["D2"].value] == ["77", 1234.6, dt.datetime(2026, 3, 4), 5]
    assert ws["A2"].number_format == "@" and ws["B2"].number_format == "0.0"
    assert ws["C2"].number_format == "dd-mmm-yyyy" and ws["D2"].number_format == "0"
    assert ws["D3"].value is None and any("'oops'" in n and "whole number" in n for n in run["notes"])   # said, not hidden
    assert ws["E2"].value == "=B2*0.1" and ws["E2"].number_format == "0.000"           # formula untouched, display format set


def test_text_lock_never_turns_a_formula_cell_into_text():
    st, s = system()
    s = lock(st, s, dict(column="Fee", type="text"))
    run, wb = output_book(st, s, book(INPUT))
    assert wb["Out"]["E2"].value == "=B2*0.1" and wb["Out"]["E2"].number_format == "General"


def test_auto_removes_the_lock_and_bad_requests_are_rejected():
    st, s = system()
    s = lock(st, s, dict(column="Amt", type="integer"))
    assert col(s, "Amt")["dtype"] == {"type": "integer"}
    s = lock(st, s, dict(column="Amt", type="auto"))
    assert "dtype" not in col(s, "Amt")
    for edit, msg in ((dict(column="Amt", type="banana"), "unknown type"), (dict(column="Nope", type="text"), "no such column"),
                      (dict(column="Amt", type="decimal", decimals=99), "decimals"),
                      (dict(column="Day", type="date", format="<script>"), "date format")):
        with pytest.raises(ValueError, match=msg):
            lock(st, s, edit)


def test_unlocked_columns_are_pasted_as_they_come():
    st, s = system()
    run, wb = output_book(st, s, book([(5, "x", "y", 1.5)]))
    assert [c.value for c in wb["Out"][2]][:4] == [5, "x", "y", 1.5]


def test_api_round_trip():
    from test_api import client, upload
    c = client()
    did = c.post("/api/drafts").get_json()["draft_id"]
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Amt"])
    o.append(["A1", 1.5])
    upload(c, f"/api/drafts/{did}/files", "i.xlsx", book([("A1", 1.5, dt.datetime(2026, 1, 1), 1)]), hint="input")
    upload(c, f"/api/drafts/{did}/files", "o.xlsx", b(out), hint="output")
    sid = c.post(f"/api/drafts/{did}/learn", json={"name": "api"}).get_json()["system"]["id"]
    r = c.put(f"/api/systems/{sid}", json={"dtypes": [{"sheet": "Out", "column": "ID", "type": "text"}]}).get_json()
    assert r["success"] and r["system"]["sheets"][0]["columns"][0]["dtype"] == {"type": "text"}
    bad = c.put(f"/api/systems/{sid}", json={"dtypes": [{"sheet": "Out", "column": "ID", "type": "xyz"}]})
    assert bad.status_code == 400 and "unknown type" in bad.get_json()["error"]


def test_day_first_option_decides_ambiguous_dates():
    cv = engine.coerce_value
    assert cv("04/05/2026", {"type": "date"})[0] == dt.datetime(2026, 4, 5)
    assert cv("04/05/2026", {"type": "date", "dayfirst": True})[0] == dt.datetime(2026, 5, 4)
    st, s = system()
    s = lock(st, s, dict(column="Day", type="date", dayfirst=True))
    assert col(s, "Day")["dtype"] == {"type": "date", "format": "yyyy-mm-dd", "dayfirst": True}
