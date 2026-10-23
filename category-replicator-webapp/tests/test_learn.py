import pandas as pd
from helpers import make_case
from replicator import engine
from replicator.tables import read_workbook


def entries(inp, mp, out):
    es = []
    for role, name, data in (("input", "in.xlsx", inp), ("mapping", "map.xlsx", mp), ("output", "out.xlsx", out)):
        for s in read_workbook(data, name):
            es.append(dict(file=name, sheet=s["sheet"], role=role, df=s["df"], header_row=s["header_row"],
                           n_formulas=s["n_formulas"], bytes=data))
    return es


def test_learns_and_replays():
    inp, mp, out, rows, out_rows = make_case(n=80, split=True)
    system = engine.build_system(entries(inp, mp, out), {"name": "t"})
    by = {r["name"]: r for r in system["rules"]}
    for n, r in by.items():
        print(n, r["kind"], r["status"], r["match"], "|", r["description"])
    assert by["ID"]["kind"] == "copy"
    assert by["Rate"]["kind"] == "lookup" and by["Rate"]["table"] == "Rates"
    assert by["Group"]["kind"] == "lookup" and by["Group"]["table"] == "Groups"
    assert by["Fee"]["kind"] == "arith"
    assert by["Net"]["kind"] == "arith"
    assert by["RWA"]["kind"] == "arith"
    assert by["Bucket"]["kind"] == "cond"
    assert by["Label"]["kind"] == "concat"
    assert by["Year"]["kind"] == "datepart"
    assert by["Name"]["kind"] == "text"
    assert all(r["status"] == "ok" for r in system["rules"]), [(r["name"], r["status"], r["match"]) for r in system["rules"]]
    sp = system["filters"]["split"]
    assert sp["enabled"] and sp["column"] == "Region"

    # replay on training data
    X = read_workbook(inp, "in.xlsx")[0]["df"]
    res = engine.apply_system(system, X)
    assert set(res["sheets"]) == {"EU", "US", "ASIA"}
    Y = pd.concat([s["df"] for s in read_workbook(out, "out.xlsx")], ignore_index=True)
    cmp_ = engine.compare_frames(res["out"], Y)
    assert cmp_["match"] == 1.0, cmp_

    # brand new data
    inp2, _, _, rows2, out_rows2 = make_case(n=40, seed=99, split=True)
    X2 = read_workbook(inp2, "in2.xlsx")[0]["df"]
    r2 = engine.apply_system(system, X2)
    exp = pd.DataFrame(out_rows2)
    assert engine.compare_frames(r2["out"], exp[system["output"]["columns"]])["match"] == 1.0


def test_workbook_and_live_formulas():
    inp, mp, out, *_ = make_case(n=30, split=False)
    system = engine.build_system(entries(inp, mp, out), {"name": "t"})
    assert not system["filters"]["split"]["enabled"]
    X = read_workbook(inp, "in.xlsx")[0]["df"]
    res = engine.apply_system(system, X)
    import openpyxl, io
    for live in (False, True):
        data = engine.write_workbook(system, res["sheets"], res["work"], live_formulas=live)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert "Run summary" in wb.sheetnames
        ws = wb["Result"]
        hdr = [c.value for c in ws[1]]
        if live:
            fee = hdr.index("Fee")
            assert str(ws.cell(row=2, column=fee + 1).value).startswith("=")
            print(ws.cell(row=2, column=fee + 1).value, ws.cell(row=2, column=hdr.index("Rate") + 1).value)
            assert "Rates" in wb.sheetnames
