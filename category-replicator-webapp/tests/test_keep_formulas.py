"""A formula in the sample is replicated as written - even one our own engine cannot calculate."""
import io
import tempfile

import openpyxl

from replicator import engine, service
from replicator.store import LocalStore

DATA = [(1, "x", 10), (2, "y", 20), (3, "x", 30), (4, "y", 5)]


def b(wb):
    x = io.BytesIO()
    wb.save(x)
    return x.getvalue()


def book(rows, header="ID,Grp,Amt"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(header.split(","))
    [ws.append(list(r)) for r in rows]
    return b(wb)


def system():
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Grp", "Amt", "GrpTotal", "Idx", "Doubled"])
    tot = {"x": 40, "y": 25}
    cached = {}
    for i, (a, g, m) in enumerate(DATA, start=2):
        o.append([a, g, m, f"=SUMIFS($C$2:$C$5,$B$2:$B$5,B{i})", f"=INDEX($C$2:$C$5,MATCH(A{i},$A$2:$A$5,0))", f"=C{i}*2"])
        cached.update({f"D{i}": tot[g], f"E{i}": m, f"F{i}": m * 2})
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    service.add_draft_files(st, did, [("i.xlsx", book(DATA))], "input")
    service.add_draft_files(st, did, [("o.xlsx", engine.inject_cached(b(out), ["Out"], {"Out": cached}))], "output")
    return st, service.create_system(st, did, {"name": "keep"})


def test_unsupported_functions_are_kept_not_guessed():
    st, s = system()
    by = {r["name"]: r for r in s["rules"]}
    assert by["GrpTotal"]["kind"] == "formula" and by["GrpTotal"]["status"] == "excel"      # not IF(Grp=="x",40,25)
    assert by["Idx"]["kind"] == "formula" and by["Idx"]["status"] == "excel"                 # not a copy of Amt
    assert by["Doubled"]["status"] == "ok"                                                   # supported: calculated too


def test_formulas_written_verbatim_even_for_unsupported_functions():
    st, s = system()
    new = [(7, "x", 1), (8, "y", 2), (9, "x", 3), (10, "x", 4), (11, "y", 5), (12, "y", 6)]
    run = service.run_system(st, s["id"], [("n.xlsx", book(new))], live_formulas=True)
    assert any("Calculated by Excel itself" in n and "GrpTotal" in n for n in run["notes"])
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    ws = openpyxl.load_workbook(io.BytesIO(data))["Out"]
    assert [c.value for c in ws[2]][3:] == ["=SUMIFS($C$2:$C$5,$B$2:$B$5,B2)", "=INDEX($C$2:$C$5,MATCH(A2,$A$2:$A$5,0))", "=C2*2"]
    assert [c.value for c in ws[7]][3:] == ["=SUMIFS($C$2:$C$5,$B$2:$B$5,B7)", "=INDEX($C$2:$C$5,MATCH(A7,$A$2:$A$5,0))", "=C7*2"]   # fill-down only
    vals = service.run_system(st, s["id"], [("n.xlsx", book(new))], live_formulas=False)
    assert any("stay EMPTY" in n for n in vals["notes"])
