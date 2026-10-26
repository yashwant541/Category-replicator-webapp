"""Generate a recipe from the dummy example, run it with a fake `dataiku` on NEW input, compare to expected."""
import io
import pathlib
import subprocess
import sys
import tempfile
import types

import openpyxl
import pytest

from replicator import codegen, service
from replicator.store import LocalStore

ROOT = pathlib.Path(__file__).resolve().parents[1]
D = ROOT / "dev" / "sample_data"


class FakeFolder:
    registry = {}

    def __init__(self, fid):
        self.files = FakeFolder.registry.setdefault(fid, {})

    def list_paths_in_partition(self):
        return list(self.files)

    def get_download_stream(self, p):
        return io.BytesIO(self.files[p])

    def upload_stream(self, p, s):
        self.files[p] = s.read()


def test_generated_recipe_runs_on_new_data(monkeypatch):
    subprocess.check_call([sys.executable, str(ROOT / "dev" / "make_dummy_data.py")], stdout=subprocess.DEVNULL)
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    for f, h in (("1_input_exposures.xlsx", "input"), ("2_mapping_tables.xlsx", "mapping"),
                 ("3_expected_output.xlsx", "output")):
        service.add_draft_files(st, did, [(f, (D / f).read_bytes())], h)
    s = service.create_system(st, did, {"name": "dummy"})
    code = codegen.generate_script(s, data_folder="IN", output_folder="OUT", template_folder="TPL", category="Dummy")

    fake = types.ModuleType("dataiku")
    fake.Folder = FakeFolder
    monkeypatch.setitem(sys.modules, "dataiku", fake)
    FakeFolder.registry.clear()
    FakeFolder("IN").files["/new.xlsx"] = (D / "4_new_input_exposures.xlsx").read_bytes()
    # Python-level import isolation: the recipe registers its own replicator modules
    saved = {k: v for k, v in sys.modules.items() if k.startswith("replicator")}
    for k in saved:
        del sys.modules[k]
    try:
        exec(compile(code, "recipe", "exec"), {"__name__": "recipe"})
    finally:
        for k in [k for k in sys.modules if k.startswith("replicator")]:
            del sys.modules[k]
        sys.modules.update(saved)

    outs = FakeFolder("OUT").files
    tpls = FakeFolder("TPL").files
    assert len(outs) == 1 and len(tpls) == 1
    got = openpyxl.load_workbook(io.BytesIO(next(iter(outs.values()))), data_only=True)
    exp = openpyxl.load_workbook(D / "5_new_expected_output.xlsx", data_only=True)
    for sh in ("Retail_EMEA", "Retail_APAC", "Corporate_Global", "SME_Global"):
        g = [[c for c in r] for r in got[sh].iter_rows(values_only=True)]
        e = [[c for c in r] for r in exp[sh].iter_rows(values_only=True)]
        assert len(g) == len(e), sh
        for gr, er in zip(g, e):
            for a, b in zip(gr, er):
                if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                    assert abs(a - b) < 1e-6 * max(1, abs(b)), (sh, a, b)
                else:
                    assert (a or "") == (b or ""), (sh, a, b)
    live = openpyxl.load_workbook(io.BytesIO(next(iter(tpls.values()))))
    assert any(str(c.value).startswith("=") for ws in live for row in ws.iter_rows() for c in row)
