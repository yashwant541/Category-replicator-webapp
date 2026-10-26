"""Role guessing of uploaded workbooks: mapping sheets inside the output file are mapping."""
import io
import openpyxl
import pandas as pd


def test_output_workbook_mapping_sheets_become_mapping():
    import io, openpyxl
    from replicator import service
    from replicator.store import LocalStore
    import tempfile
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Tier", "Amount", "Weight"])
    for r in range(2, 5):
        o.append([r, "A", 10 * r, f"=VLOOKUP(B{r},Weights!$A$2:$B$3,2,FALSE)"])
    w = out.create_sheet("Weights")
    w.append(["Tier", "W"])
    w.append(["A", 0.5])
    w.append(["B", 1])
    bio = io.BytesIO(); out.save(bio)
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    d = service.add_draft_files(st, did, [("out.xlsx", bio.getvalue())], "output")
    roles = {s["sheet"]: s["role"] for s in d["files"][0]["sheets"]}
    assert roles == {"Out": "output", "Weights": "mapping"}


def test_value_only_mapping_sheet_next_to_formula_sheets_is_mapping():
    import io, tempfile, openpyxl
    from replicator import service
    from replicator.store import LocalStore
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Result"
    o.append(["ID", "Tier", "W"])
    o.append([1, "A", "=VLOOKUP(B2,Rates!$A$2:$B$3,2,FALSE)"])
    r = out.create_sheet("Rates"); r.append(["Tier", "W"]); r.append(["A", 1]); r.append(["B", 2])
    ex = out.create_sheet("Lookup_Extra"); ex.append(["k", "v"]); ex.append(["a", 1])   # not referenced by name
    bio = io.BytesIO(); out.save(bio)
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    d = service.add_draft_files(st, did, [("o.xlsx", bio.getvalue())], "output")
    roles = {s["sheet"]: s["role"] for s in d["files"][0]["sheets"]}
    assert roles == {"Result": "output", "Rates": "mapping", "Lookup_Extra": "mapping"}, roles
