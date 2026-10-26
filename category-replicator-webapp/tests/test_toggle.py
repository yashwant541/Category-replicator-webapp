"""A Toggle sheet (Label | Value) in the mapping and in the output sample drives derived-column formulas."""
import io
import pathlib
import tempfile

import openpyxl
import pandas as pd
import pytest

from replicator import engine, service
from replicator.store import LocalStore

TOGGLE_FILE = pathlib.Path(__file__).resolve().parents[1] / "dev" / "sample_data" / "toggle" / "Toggle_sample.xlsx"
RATES = [("A", 0.10), ("B", 0.20), ("C", 0.30)]
ROWS = [(1, "A", 1000.0), (2, "B", 2500.0), (3, "C", 400.0), (4, "A", 150.0), (5, "B", 999.0)]


def b(wb):
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def input_wb(rows=ROWS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["ID", "Tier", "Amount"])
    for r in rows:
        ws.append(list(r))
    return b(wb)


def mapping_wb(indicator=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["Tier", "Rate"])
    for r in RATES:
        ws.append(list(r))
    t = wb.create_sheet("Toggle")
    t.append(["Label", "Value"])
    t.append(["Indicator", indicator])
    return b(wb)


def output_wb(indicator=True):
    """Output sample: Fee depends on the toggle on the Toggle sheet; cached results saved like Excel does."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append(["ID", "Tier", "Amount", "Rate", "Fee"])
    cached = {}
    rate = dict(RATES)
    for i, (a, t, amt) in enumerate(ROWS, start=2):
        ws.append([a, t, amt, f"=VLOOKUP(B{i},Rates!$A$2:$B$4,2,FALSE)", f"=IF(Toggle!$B$2,ROUND(C{i}*D{i},2),0)"])
        cached[f"D{i}"] = rate[t]
        cached[f"E{i}"] = round(amt * rate[t], 2) if indicator else 0
    r = wb.create_sheet("Rates")
    r.append(["Tier", "Rate"])
    for x in RATES:
        r.append(list(x))
    tg = wb.create_sheet("Toggle")
    tg.append(["Label", "Value"])
    tg.append(["Indicator", indicator])
    return engine.inject_cached(b(wb), ["Result"], {"Result": cached})


def build(indicator=True):
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    d = None
    for n, data, h in (("in.xlsx", input_wb(), "input"), ("map.xlsx", mapping_wb(indicator), "mapping"),
                       ("out.xlsx", output_wb(indicator), "output")):
        d = service.add_draft_files(st, did, [(n, data)], h)
    roles = {s["sheet"]: s["role"] for s in d["files"][2]["sheets"]}
    assert roles == {"Result": "output", "Rates": "mapping", "Toggle": "mapping"}, roles
    return st, service.create_system(st, did, {"name": "tog"})


def rows_of(run, sheet="Result"):
    p = run["preview"][sheet]
    return [dict(zip(p["headers"], r)) for r in p["data"]]


def test_toggle_learned_and_used():
    st, s = build(True)
    assert [t["name"] for t in s["toggles"]] == ["Indicator"] and s["toggles"][0]["value"] == "Yes"
    assert s["toggles"][0]["sheet"] == "Toggle" and s["toggles"][0]["cell"] == "B2"
    assert all(r["status"] == "ok" for r in s["rules"]), [(r["name"], r["status"]) for r in s["rules"]]
    assert set(s["tables"]) == {"Rates", "Toggle"}
    run = service.run_system(st, s["id"], [("new.xlsx", input_wb([(9, "B", 500.0), (10, "C", 100.0)]))])
    assert [r["Fee"] for r in rows_of(run)] == [100.0, 30.0]
    assert run["toggles"][0]["value"] == "Yes"


def test_mapping_toggle_value_wins_at_run_time():
    """Attach a mapping whose Toggle sheet says FALSE: formulas must use it (not the sample's TRUE)."""
    st, s = build(True)
    run = service.run_system(st, s["id"], [("new.xlsx", input_wb([(9, "B", 500.0)]))],
                             mapping_files=[("map2.xlsx", mapping_wb(False))])
    assert rows_of(run)[0]["Fee"] == 0 and run["toggles"][0]["value"] == "No"


def test_per_run_override_and_written_formulas():
    st, s = build(True)
    inp = input_wb([(9, "B", 500.0)])
    run = service.run_system(st, s["id"], [("new.xlsx", inp)], live_formulas=True, toggle_values={"Indicator": "No"})
    assert rows_of(run)[0]["Fee"] == 0
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    f = openpyxl.load_workbook(io.BytesIO(data))
    ws = f["Result"]
    hdr = [c.value for c in ws[1]]
    fee = ws.cell(2, hdr.index("Fee") + 1).value
    assert fee == "=IF(Toggle!$B$2,ROUND(C2*D2,2),0)", fee            # the toggle is referenced, not baked in
    assert "Toggle" in f.sheetnames and f["Toggle"]["B2"].value is False   # sheet carries the value used
    run = service.run_system(st, s["id"], [("new.xlsx", inp)], live_formulas=True)
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    assert openpyxl.load_workbook(io.BytesIO(data))["Toggle"]["B2"].value is True


def test_users_toggle_sample_file_is_read():
    """The user's own Toggle_sample.xlsx: Label | Value, Indicator = TRUE."""
    from replicator.tables import read_workbook
    sheets = read_workbook(TOGGLE_FILE.read_bytes(), TOGGLE_FILE.name)
    assert [s["sheet"] for s in sheets] == ["Toggle"]
    assert sheets[0]["grid"] == [["Label", "Value"], ["Indicator", True]]
    assert list(sheets[0]["df"].columns) == ["Label", "Value"] and len(sheets[0]["df"]) == 1


def test_toggle_sheet_without_header_row():
    from replicator.tables import read_workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Toggle"
    ws.append(["Include Tax", "Yes"])
    ws.append(["Round Up", "No"])
    s = read_workbook(b(wb), "t.xlsx")[0]
    assert len(s["df"]) == 2 and s["grid"][0] == ["Include Tax", "Yes"]
