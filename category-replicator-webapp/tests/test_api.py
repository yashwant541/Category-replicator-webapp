import io
import os
import pathlib
import tempfile

import openpyxl
from flask import Flask
from helpers import make_case


def client():
    os.environ["REPLICATOR_LOCAL_ROOT"] = tempfile.mkdtemp()
    app = Flask(__name__)
    g = {"app": app, "__name__": "b"}
    src = (pathlib.Path(__file__).resolve().parents[1] / "backend" / "backend.py").read_text(encoding="utf-8")
    exec(compile(src, "backend.py", "exec"), g)
    return app.test_client()


def upload(c, path, name, data, **form):
    return c.post(path, data={"files": (io.BytesIO(data), name), **form},
                  content_type="multipart/form-data").get_json()


def test_full_flow():
    c = client()
    inp, mp, out, *_ = make_case(n=70)
    did = c.post("/api/drafts").get_json()["draft_id"]
    for name, data, hint in (("in.xlsx", inp, "input"), ("rates_map.xlsx", mp, "mapping"), ("out.xlsx", out, "output")):
        r = upload(c, f"/api/drafts/{did}/files", name, data, hint=hint)
        assert r["success"], r
    d = r["draft"]
    roles = {s["sheet"]: s["role"] for f in d["files"] for s in f["sheets"] if f["name"] == "out.xlsx"}
    assert set(roles.values()) == {"output"}
    r = c.post(f"/api/drafts/{did}/learn", json={"name": "Demo cat"}).get_json()
    assert r["success"], r
    s = r["system"]
    sid = s["id"]
    assert c.get("/api/systems").get_json()["systems"][0]["name"] == "Demo cat"
    inp2, *_ = make_case(n=30, seed=5)
    r = upload(c, f"/api/systems/{sid}/run", "new.xlsx", inp2)
    assert r["success"], r
    run = r["run"]
    assert set(run["sheets"]) == {"EU", "US", "ASIA"} and run["rows"] == 30
    dl = c.get(f"/api/systems/{sid}/runs/{run['run_id']}/download")
    wb = openpyxl.load_workbook(io.BytesIO(dl.data))
    assert {"EU", "US", "ASIA", "Run summary"} <= set(wb.sheetnames)
    # turn the split off, round trip
    s["filters"]["split"]["enabled"] = False
    assert c.put(f"/api/systems/{sid}", json={"filters": s["filters"]}).get_json()["success"]
    r = upload(c, f"/api/systems/{sid}/run", "new.xlsx", inp2)
    assert len(r["run"]["sheets"]) == 1
    # wrong input
    bad = openpyxl.Workbook()
    bad.active.append(["x", "y"])
    bad.active.append([1, 2])
    bio = io.BytesIO()
    bad.save(bio)
    r = upload(c, f"/api/systems/{sid}/run", "bad.xlsx", bio.getvalue())
    assert not r["success"] and "no sheet looks like the input" in r["error"]


def test_uncached_formula_mapping_sheet_is_reported():
    c = client()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Good"
    ws.append(["k", "v"]); ws.append(["x", 1])
    bad = wb.create_sheet("Calc")
    bad.append(["k", "v"]); bad.append(["x", "=1+1"]); bad.append(["y", "=2*3"])
    allf = wb.create_sheet("AllFormulas")
    allf.append(['="k"', '="v"']); allf.append(['="x"', "=1"])
    bio = io.BytesIO(); wb.save(bio)
    did = c.post("/api/drafts").get_json()["draft_id"]
    r = upload(c, f"/api/drafts/{did}/files", "map.xlsx", bio.getvalue(), hint="mapping")
    assert r["success"]
    issues = {i["sheet"]: i["issue"] for i in r["draft"]["files"][0]["issues"]}
    assert "no saved result" in issues["Calc"] and "skipped" in issues["AllFormulas"]
    assert [s["sheet"] for s in r["draft"]["files"][0]["sheets"]] == ["Good", "Calc"]


def test_overlapping_uploads_do_not_lose_files():
    import threading
    c = client()
    inp, mp, out, *_ = make_case(n=40)
    did = c.post("/api/drafts").get_json()["draft_id"]
    errs = []

    def go(name, data, hint):
        r = upload(c, f"/api/drafts/{did}/files", name, data, hint=hint)
        if not r.get("success"):
            errs.append(r)
    ts = [threading.Thread(target=go, args=a) for a in (("in.xlsx", inp, "input"), ("map.xlsx", mp, "mapping"), ("out.xlsx", out, "output"))]
    [t.start() for t in ts]; [t.join() for t in ts]
    d = c.get(f"/api/drafts/{did}").get_json()["draft"]
    assert not errs and sorted(f["name"] for f in d["files"]) == ["in.xlsx", "map.xlsx", "out.xlsx"]
