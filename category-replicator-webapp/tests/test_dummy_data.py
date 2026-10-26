"""End-to-end on dev/sample_data: 4 input sheets, 4 mapping sheets, 4 output sheets with different
derived columns + mapping sheets inside the output workbook."""
import io
import pathlib
import subprocess
import sys
import tempfile

import openpyxl
import pytest

from replicator import service
from replicator.store import LocalStore

ROOT = pathlib.Path(__file__).resolve().parents[1]
D = ROOT / "dev" / "sample_data"


@pytest.fixture(scope="module")
def system():
    subprocess.check_call([sys.executable, str(ROOT / "dev" / "make_dummy_data.py")], stdout=subprocess.DEVNULL)
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    for f, hint in (("1_input_exposures.xlsx", "input"), ("2_mapping_tables.xlsx", "mapping"),
                    ("3_expected_output.xlsx", "output")):
        service.add_draft_files(st, did, [(f, (D / f).read_bytes())], hint)
    return st, service.create_system(st, did, {"name": "dummy"})


def test_learned_from_formulas(system):
    st, s = system
    assert all(r["status"] == "ok" for r in s["rules"]), [(r["name"], r["status"]) for r in s["rules"]]
    assert s["training"]["replay_match"] == 1.0
    assert set(s["tables"]) == {"Ratings", "Products", "Countries", "FX"}
    sc = s["output"]["sheet_columns"]
    assert sc["Retail_EMEA"][-1] == "Expected_Loss" and sc["SME_Global"][-1] == "Label"
    assert "Net_Exposure" in sc["Corporate_Global"] and "Net_Exposure" not in sc["SME_Global"]
    assert s["filters"]["split"]["column"] == "Source sheet"
    raw = {"LoanID", "Customer", "Segment", "Region", "Product", "Rating", "Country", "Currency",
           "Exposure", "Collateral", "Maturity"}
    assert set(s["input"]["required"]) <= raw | {"Source sheet"}          # derived columns are never required


def test_new_input_matches_expected(system):
    st, s = system
    exp_wb = openpyxl.load_workbook(D / "5_new_expected_output.xlsx", data_only=True)
    run = service.run_system(st, s["id"], [("new.xlsx", (D / "4_new_input_exposures.xlsx").read_bytes())],
                             expected_file=("exp.xlsx", (D / "5_new_expected_output.xlsx").read_bytes()))
    assert set(run["sheets"]) == {"Retail_EMEA", "Retail_APAC", "Corporate_Global", "SME_Global"}
    for sh in run["sheets"]:
        want = [c.value for c in exp_wb[sh][1]]
        assert run["preview"][sh]["headers"] == want, (sh, run["preview"][sh]["headers"], want)
        assert run["sheets"][sh] == exp_wb[sh].max_row - 1
    assert run["compare"]["match"] == 1.0, run["compare"]
    assert not run["rule_errors"]
    assert not [n for n in run["notes"] if "no match" in n], run["notes"]    # no spurious mapping warnings
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames[:4] == ["Retail_EMEA", "Retail_APAC", "Corporate_Global", "SME_Global"]


def test_live_formulas_per_sheet_layout(system):
    st, s = system
    run = service.run_system(st, s["id"], [("new.xlsx", (D / "4_new_input_exposures.xlsx").read_bytes())],
                             live_formulas=True)
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["SME_Global"]
    hdr = [c.value for c in ws[1]]
    assert str(ws.cell(row=2, column=hdr.index("Capital_Required") + 1).value).startswith("=")
    assert {"Ratings", "FX"} <= set(wb.sheetnames)


def test_formulas_written_into_output_cells(system):
    """The generated workbook carries the formulas (as in the example), with saved results too."""
    st, s = system
    run = service.run_system(st, s["id"], [("new.xlsx", (D / "4_new_input_exposures.xlsx").read_bytes())],
                             live_formulas=True)
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    f = openpyxl.load_workbook(io.BytesIO(data))
    v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ex = openpyxl.load_workbook(D / "5_new_expected_output.xlsx")
    exv = openpyxl.load_workbook(D / "5_new_expected_output.xlsx", data_only=True)
    n_formula = 0
    for sh in ("Retail_EMEA", "Retail_APAC", "Corporate_Global", "SME_Global"):
        hdr = [c.value for c in f[sh][1]]
        base = 11
        for j, col in enumerate(hdr[base:], start=base + 1):
            got, want = f[sh].cell(2, j).value, ex[sh].cell(2, j).value
            assert got == want, (sh, col, got, want)           # same formula text as the example sheet
            n_formula += 1
            gv, wv = v[sh].cell(2, j).value, exv[sh].cell(2, j).value
            assert gv == wv or abs(float(gv) - float(wv)) < 1e-6, (sh, col, gv, wv)   # saved result present
    assert n_formula >= 25
    assert {"Ratings", "Products", "Countries", "FX"} <= set(f.sheetnames)
