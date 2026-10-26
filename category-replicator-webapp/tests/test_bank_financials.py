"""Real-world style case: one input sheet, a mapping that routes rows to output sheets (derived column),
an audit sheet repeating every row, a column-vs-column IF and no formulas at all (values only)."""
import io
import pathlib
import random
import tempfile

import openpyxl
import pandas as pd
import pytest

from replicator import service
from replicator.store import LocalStore

B = pathlib.Path(__file__).resolve().parents[1] / "dev" / "sample_data" / "bank"


@pytest.fixture(scope="module")
def system():
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    for f, hint in (("bank_financials_source_data.xlsx", "input"), ("bank_financials_mapping.xlsx", "mapping"),
                    ("bank_financials_outputs.xlsx", "output")):
        d = service.add_draft_files(st, did, [(f, (B / f).read_bytes())], hint)
    roles = {s["sheet"]: s["role"] for f in service.get_draft(st, did)["files"] for s in f["sheets"]}
    assert roles["Data_Dictionary"] == "ignore"
    return st, service.create_system(st, did, {"name": "bank"})


def test_learns(system):
    st, s = system
    assert all(r["status"] == "ok" for r in s["rules"]), [(r["name"], r["status"]) for r in s["rules"]]
    assert s["training"]["replay_match"] == 1.0
    sheets = {x["name"]: x["conditions"] for x in s["filters"]["split"]["sheets"]}
    assert sheets["All_Enriched_Audit"] == []
    assert sheets["Vendor_Output"][0]["col"] == "Output_Sheet"


def test_new_data(system):
    st, s = system
    src = pd.read_excel(B / "bank_financials_source_data.xlsx", sheet_name="Raw_Data")
    mp = pd.read_excel(B / "bank_financials_mapping.xlsx").set_index("Account_Type")
    rnd = random.Random(5)
    new = src.sample(45, random_state=3, replace=True).reset_index(drop=True)
    new["Transaction_ID"] = ["TXN9%04d" % i for i in range(len(new))]
    new["Amount_INR"] = [round(rnd.uniform(1000, 450000), 2) for _ in range(len(new))]
    new["Signed_Amount_INR"] = [a if d == "Credit" else -a for a, d in zip(new.Amount_INR, new.Debit_Credit)]
    new["Transaction_Date"] = pd.to_datetime(new["Transaction_Date"]) + pd.to_timedelta(rnd.randint(1, 60), unit="D")
    bio = io.BytesIO()
    new.to_excel(bio, sheet_name="Raw_Data", index=False)
    run = service.run_system(st, s["id"], [("new.xlsx", bio.getvalue())])
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    out = pd.read_excel(io.BytesIO(data), sheet_name=None)
    # independent expectation
    exp = new.copy()
    for c in ("Output_Sheet", "Derived_Category", "GL_Code", "Review_Threshold_INR", "Owner_Team"):
        exp[c] = exp.Account_Type.map(mp[c])
    exp["Review_Flag"] = (exp.Amount_INR > exp.Review_Threshold_INR).map({True: "REVIEW", False: "OK"})
    exp["Month"] = pd.to_datetime(exp.Transaction_Date).dt.strftime("%Y-%m")
    exp["Net_Amount_INR"] = exp.Signed_Amount_INR
    assert len(out["All_Enriched_Audit"]) == len(new)
    for sh in ("Payroll_Output", "Vendor_Output", "Customer_Output"):
        e = exp[exp.Output_Sheet == sh]
        assert len(out[sh]) == len(e), sh
        assert list(out[sh].Transaction_ID) == list(e.Transaction_ID)
        assert list(out[sh].Review_Flag) == list(e.Review_Flag)
        assert list(out[sh].GL_Code) == list(e.GL_Code)
        assert list(out[sh].Month) == list(e.Month)
        assert "Output_Sheet" not in out[sh].columns            # only the audit sheet has it
    assert list(out["All_Enriched_Audit"].Output_Sheet) == list(exp.Output_Sheet)
    assert not run["rule_errors"]


def test_run_with_expected_output_compares_per_sheet(system):
    st, s = system
    run = service.run_system(st, s["id"], [("src.xlsx", (B / "bank_financials_source_data.xlsx").read_bytes())],
                             expected_file=("out.xlsx", (B / "bank_financials_outputs.xlsx").read_bytes()))
    assert run["compare"]["aligned"] and run["compare"]["match"] == 1.0, run["compare"]
    assert run["sheets"] == {"Payroll_Output": 4, "Vendor_Output": 15, "Customer_Output": 41, "All_Enriched_Audit": 60}


def test_live_formulas_no_self_reference_and_cached(system):
    st, s = system
    run = service.run_system(st, s["id"], [("src.xlsx", (B / "bank_financials_source_data.xlsx").read_bytes())],
                             live_formulas=True)
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    f = openpyxl.load_workbook(io.BytesIO(data))
    v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws, wv = f["Vendor_Output"], v["Vendor_Output"]
    hdr = [c.value for c in ws[1]]
    assert ws.cell(2, 1).value == "TXN00001"                         # raw column is a value, not =A2
    j = hdr.index("Review_Flag") + 1
    assert str(ws.cell(2, j).value).startswith("=IF(") and wv.cell(2, j).value in ("OK", "REVIEW")
    j = hdr.index("GL_Code") + 1
    assert "INDEX(Mapping!" in ws.cell(2, j).value and wv.cell(2, j).value == "GL-6200"
    assert "Mapping" in f.sheetnames
