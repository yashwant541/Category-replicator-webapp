"""Viewing / editing / saving the stored formulas of a system."""
import io
import pathlib
import tempfile

import openpyxl
import pytest

from replicator import service
from replicator.store import LocalStore
from test_template import B, D, b, make, output_book, small_book

ROWS = [(1, "x", 10), (2, "y", 20), (3, "x", 30)]


def system():
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Grp", "Amt", "Tax", "Note"])
    for i, (a, g, m) in enumerate(ROWS, start=2):
        o.append([a, g, m, f"=C{i}*0.1", None])
    return make([("i.xlsx", small_book(ROWS), "input"), ("o.xlsx", b(out), "output")])


def col(s, name):
    return next(c for c in s["sheets"][0]["columns"] if c["name"] == name)


def test_edit_a_formula_reference_and_use_it_on_the_next_run():
    st, s = system()
    assert col(s, "Tax")["raw"] == "=C2*0.1" and col(s, "Note")["kind"] == "empty"
    s = service.update_system(st, s["id"], {"formulas": [{"sheet": "Out", "column": "Tax", "raw": "=C2*0.2+$H$1"}]})
    assert col(s, "Tax")["raw"] == "=C2*0.2+$H$1" and col(s, "Tax")["orig"] == "=C2*0.1"      # original remembered
    s = service.get_system(st, s["id"])                                                         # and it was saved
    assert col(s, "Tax")["raw"] == "=C2*0.2+$H$1"
    run, wb = output_book(st, s, small_book([(7, "q", 5), (8, "r", 6)]))
    assert [wb["Out"]["D2"].value, wb["Out"]["D3"].value] == ["=C2*0.2+$H$1", "=C3*0.2+$H$1"]


def test_reset_by_typing_the_original_back_removes_the_edited_mark():
    st, s = system()
    s = service.update_system(st, s["id"], {"formulas": [{"sheet": "Out", "column": "Tax", "raw": "=C2*9"}]})
    assert "orig" in col(s, "Tax")
    s = service.update_system(st, s["id"], {"formulas": [{"sheet": "Out", "column": "Tax", "raw": "=C2*0.1"}]})
    assert "orig" not in col(s, "Tax") and col(s, "Tax")["raw"] == "=C2*0.1"


def test_give_an_empty_column_a_formula_and_clear_one():
    st, s = system()
    s = service.update_system(st, s["id"], {"formulas": [{"sheet": "Out", "column": "Note", "raw": '=IF(C2>15,"big","small")'},
                                                         {"sheet": "Out", "column": "Tax", "raw": ""}]})
    assert col(s, "Note")["kind"] == "formula" and col(s, "Note")["sample_row"] == 2
    assert col(s, "Tax")["kind"] == "empty"
    run, wb = output_book(st, s, small_book([(7, "q", 5), (8, "r", 6)]))
    ws = wb["Out"]
    assert ws["E3"].value == '=IF(C3>15,"big","small")' and ws["D2"].value is None


def test_bad_formulas_and_input_columns_are_rejected_and_nothing_is_saved():
    st, s = system()
    for raw, msg in (("C2*2", "must start with ="), ("=SUM(C2", "not closed"), ("=C2)", "no matching"), ('=IF(C2,"a)', "quote")):
        with pytest.raises(ValueError, match=msg):
            service.update_system(st, s["id"], {"formulas": [{"sheet": "Out", "column": "Tax", "raw": raw}]})
    with pytest.raises(ValueError, match="input column"):
        service.update_system(st, s["id"], {"formulas": [{"sheet": "Out", "column": "Amt", "raw": "=1"}]})
    with pytest.raises(ValueError, match="no such column"):
        service.update_system(st, s["id"], {"formulas": [{"sheet": "Nope", "column": "Tax", "raw": "=1"}]})
    assert col(service.get_system(st, s["id"]), "Tax")["raw"] == "=C2*0.1"


def test_edit_through_the_api():
    from test_api import client, upload
    c = client()
    did = c.post("/api/drafts").get_json()["draft_id"]
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Grp", "Amt", "Tax"])
    for i, (a, g, m) in enumerate(ROWS, start=2):
        o.append([a, g, m, f"=C{i}*0.1"])
    upload(c, f"/api/drafts/{did}/files", "i.xlsx", small_book(ROWS), hint="input")
    upload(c, f"/api/drafts/{did}/files", "o.xlsx", b(out), hint="output")
    sid = c.post(f"/api/drafts/{did}/learn", json={"name": "api"}).get_json()["system"]["id"]
    r = c.put(f"/api/systems/{sid}", json={"formulas": [{"sheet": "Out", "column": "Tax", "raw": "=C2*3"}]}).get_json()
    assert r["success"] and r["system"]["sheets"][0]["columns"][-1]["raw"] == "=C2*3"
    bad = c.put(f"/api/systems/{sid}", json={"formulas": [{"sheet": "Out", "column": "Tax", "raw": "nope"}]})
    assert bad.status_code == 400 and "must start with =" in bad.get_json()["error"]
