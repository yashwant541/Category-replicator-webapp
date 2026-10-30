"""The template model: input columns pasted per sheet filter + the sample's formulas filled down. Nothing computed."""
import io
import pathlib
import tempfile

import openpyxl
import pytest

from replicator import engine, service
from replicator.store import LocalStore
from replicator.tables import read_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
D = ROOT / "dev" / "sample_data"
B = D / "bank"


def b(wb):
    x = io.BytesIO()
    wb.save(x)
    return x.getvalue()


def make(files):
    """files: [(name, bytes, role)] -> (store, system)"""
    st = LocalStore(tempfile.mkdtemp())
    did = service.new_draft(st)
    for n, data, role in files:
        service.add_draft_files(st, did, [(n, data)], role)
    return st, service.create_system(st, did, {"name": "t"})


def output_book(st, s, input_bytes, **kw):
    run = service.run_system(st, s["id"], [("new.xlsx", input_bytes)], **kw)
    data, _ = service.get_run_file(st, s["id"], run["run_id"])
    return run, openpyxl.load_workbook(io.BytesIO(data))


# ------------------------------------------------------------------ the dummy data: identical, cell by cell
@pytest.fixture(scope="module")
def dummy():
    return make([(f, (D / f).read_bytes(), r) for f, r in (("1_input_exposures.xlsx", "input"),
                                                          ("2_mapping_tables.xlsx", "mapping"),
                                                          ("3_expected_output.xlsx", "output"))])


def test_new_input_reproduces_the_expected_workbook_cell_by_cell(dummy):
    st, s = dummy
    assert s["filters"]["split"]["enabled"] and s["filters"]["split"]["column"] == "Source sheet"
    run, wb = output_book(st, s, (D / "4_new_input_exposures.xlsx").read_bytes())
    exp = openpyxl.load_workbook(D / "5_new_expected_output.xlsx")
    assert [n for n in wb.sheetnames if n != "Run summary"] == exp.sheetnames       # data sheets, then mapping sheets
    cells = diff = 0
    for sh in exp.sheetnames:
        got, want = wb[sh], exp[sh]
        assert (got.max_row, got.max_column) == (want.max_row, want.max_column), sh
        for ra, rb in zip(got.iter_rows(values_only=True), want.iter_rows(values_only=True)):
            for x, y in zip(ra, rb):
                cells += 1
                diff += x != y
    assert cells > 900 and diff == 0


def test_each_sheet_keeps_its_own_columns_and_formulas(dummy):
    st, s = dummy
    by = {d["name"]: d for d in s["sheets"]}
    last = by["Retail_EMEA"]["columns"][-1]
    assert {k: last[k] for k in ("name", "col", "kind", "raw", "sample_row")} == {
        "name": "Expected_Loss", "col": 15, "kind": "formula", "raw": "=ROUND(M2*N2*O2,2)", "sample_row": 2}
    assert by["SME_Global"]["columns"][-1]["name"] == "Label"
    raw = {d["name"]: next(c["raw"] for c in d["columns"] if c["name"] == "RWA")
           for d in s["sheets"] if d["name"] != "Retail_EMEA"}
    assert len(set(raw.values())) == 3                      # three different RWA formulas, each kept as it was


# ------------------------------------------------------------------------- the bank files (values only)
def test_bank_files_filters_from_input_columns_and_empty_derived_columns():
    st, s = make([("bank_financials_source_data.xlsx", (B / "bank_financials_source_data.xlsx").read_bytes(), "input"),
                  ("bank_financials_mapping.xlsx", (B / "bank_financials_mapping.xlsx").read_bytes(), "mapping"),
                  ("bank_financials_outputs.xlsx", (B / "bank_financials_outputs.xlsx").read_bytes(), "output")])
    sheets = {x["name"]: x["conditions"] for x in s["filters"]["split"]["sheets"]}
    assert sheets["All_Enriched_Audit"] == [] and sheets["Payroll_Output"][0]["col"] == "Account_Type"
    assert all(c["col"] in s["input"]["columns"] for conds in sheets.values() for c in conds)      # input columns only
    assert any("no formula in the sample" in w for w in s["warnings"])
    run, wb = output_book(st, s, (B / "bank_financials_source_data.xlsx").read_bytes())
    assert run["sheets"] == {"Payroll_Output": 4, "Vendor_Output": 15, "Customer_Output": 41, "All_Enriched_Audit": 60}
    assert any("No formula was found" in n for n in run["notes"])
    ws = wb["Vendor_Output"]
    hdr = [c.value for c in ws[1]]
    assert hdr[:3] == ["Transaction_ID", "Transaction_Date", "Employee_ID"] and "Review_Flag" in hdr
    assert ws.cell(2, hdr.index("Review_Flag") + 1).value is None                   # no formula in the sample -> empty
    assert "Mapping" in wb.sheetnames


# --------------------------------------------------------------------- formulas verbatim, whatever they are
DATA = [(1, "x", 10), (2, "y", 20), (3, "x", 30), (4, "y", 5)]


def small_book(rows, header=("ID", "Grp", "Amt")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    return b(wb)


def test_any_formula_is_kept_verbatim_and_only_the_row_moves():
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["ID", "Grp", "Amt", "GrpTotal", "Idx", "Prev"])
    for i, (a, g, m) in enumerate(DATA, start=2):
        o.append([a, g, m, f"=SUMIFS($C$2:$C$5,$B$2:$B$5,B{i})",
                  f"=INDEX($C$2:$C$5,MATCH(A{i},$A$2:$A$5,0))", f"=C{i}-C{max(i - 1, 1)}"])
    st, s = make([("i.xlsx", small_book(DATA), "input"), ("o.xlsx", b(out), "output")])
    assert [c["kind"] for c in s["sheets"][0]["columns"]] == ["input"] * 3 + ["formula"] * 3
    new = [(7, "x", 1), (8, "y", 2), (9, "x", 3), (10, "x", 4), (11, "y", 5), (12, "y", 6)]
    run, wb = output_book(st, s, small_book(new))
    ws = wb["Out"]
    assert [c.value for c in ws[2]][3:5] == ["=SUMIFS($C$2:$C$5,$B$2:$B$5,B2)", "=INDEX($C$2:$C$5,MATCH(A2,$A$2:$A$5,0))"]
    assert [c.value for c in ws[7]][3:5] == ["=SUMIFS($C$2:$C$5,$B$2:$B$5,B7)", "=INDEX($C$2:$C$5,MATCH(A7,$A$2:$A$5,0))"]
    assert ws["F7"].value == "=C7-C6"                                  # previous-row reference follows (fill-down)
    assert [ws.cell(r, 3).value for r in range(2, 8)] == [1, 2, 3, 4, 5, 6]


def test_shift_formula_is_excel_fill_down():
    sh = engine.shift_formula
    assert sh("=VLOOKUP(D2,Rates!$A$2:$B$3,2,FALSE)", 2, 5) == "=VLOOKUP(D5,Rates!$A$2:$B$3,2,FALSE)"
    assert sh("=C2-C1", 2, 4) == "=C4-C3"
    assert sh("=IF(A2=\"A2\",'My Sheet A1'!B2,LOG10(A2))", 2, 3) == "=IF(A3=\"A2\",'My Sheet A1'!B3,LOG10(A3))"
    assert sh("=SUM($C$2:$C$40)+Toggle!$B$2", 2, 9) == "=SUM($C$2:$C$40)+Toggle!$B$2"
    assert sh("=A2", 2, 2) == "=A2"


# ------------------------------------------------- the sheet is laid out exactly like the sample (cells, params)
def test_same_cells_as_the_sample_incl_title_row_gap_column_and_parameter_cell():
    inp = openpyxl.Workbook()
    ws = inp.active
    ws.title = "Data"
    ws.append(["ID", "Amount", "Note"])
    ws.append([1, 100, "a"])
    ws.append([2, 250, "b"])
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o["A1"] = "Report"
    o["C4"], o["D4"], o["F4"] = "ID", "Amount", "Tax"                      # title row, starts in C, gap column E
    for r, (i, a) in enumerate([(1, 100), (2, 250)], start=5):
        o.cell(r, 3, i)
        o.cell(r, 4, a)
        o.cell(r, 6, f"=D{r}*$H$1")
    o["H1"] = 0.1
    st, s = make([("i.xlsx", b(inp), "input"), ("o.xlsx", b(out), "output")])
    d = s["sheets"][0]
    assert d["origin"] == {"row": 3, "col": 2} and [c["col"] for c in d["columns"]] == [2, 3, 5]
    assert d["params"][0]["cell"] == "H1" and d["params"][0]["value"] == 0.1
    new = openpyxl.Workbook()
    n = new.active
    n.title = "Data"
    n.append(["ID", "Amount", "Note"])
    n.append([7, 40, "z"])
    run, wb = output_book(st, s, b(new))
    w = wb["Out"]
    assert [w["C4"].value, w["D4"].value, w["F4"].value] == ["ID", "Amount", "Tax"]
    assert [w["C5"].value, w["D5"].value, w["F5"].value] == [7, 40, "=D5*$H$1"] and w["E5"].value is None
    assert w["H1"].value == 0.1 and w["B5"].value is None              # parameter cell restored; "Note" not in the sample


def test_only_the_samples_input_columns_are_pasted():
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Out"
    o.append(["Amt", "ID", "Double"])                                    # subset + different order, then a formula
    for i, (a, g, m) in enumerate(DATA, start=2):
        o.append([m, a, f"=A{i}*2"])
    st, s = make([("i.xlsx", small_book(DATA), "input"), ("o.xlsx", b(out), "output")])
    run, wb = output_book(st, s, small_book([(9, "q", 7)]))
    assert [[c.value for c in r] for r in wb["Out"].iter_rows()] == [["Amt", "ID", "Double"], [7, 9, "=A2*2"]]


# ----------------------------------------------------------------------------------------- filters
def test_filters_pre_filter_and_column_split_with_a_new_value():
    amts = [500, 100, 700, 200, 300, 900, 150, 450]            # not a function of ID, so "Amt >= 300" is the only rule
    rows = [(i, "EU" if i % 2 else "US", amts[i - 1]) for i in range(1, 9)]
    inp = small_book(rows, ("ID", "Region", "Amt"))
    out = openpyxl.Workbook()
    out.remove(out.active)
    for reg in ("EU", "US"):
        w = out.create_sheet(reg)
        w.append(["ID", "Region", "Amt", "Fee"])
        for i, r in enumerate([x for x in rows if x[1] == reg and x[2] >= 300], start=2):
            w.append([r[0], r[1], r[2], f"=C{i}*0.1"])
    st, s = make([("i.xlsx", inp, "input"), ("o.xlsx", b(out), "output")])
    assert s["filters"]["pre"] == [{"col": "Amt", "op": ">=", "value": 300}]
    assert s["filters"]["split"]["column"] == "Region"
    run, wb = output_book(st, s, small_book([(1, "EU", 500), (2, "APAC", 900), (3, "US", 100)], ("ID", "Region", "Amt")))
    assert run["sheets"] == {"EU": 1, "APAC": 1}                         # 100 filtered out; APAC is a new sheet
    assert wb["APAC"]["D2"].value == "=C2*0.1" and any("new" in n for n in run["notes"])


def test_wrong_input_is_rejected_clearly():
    st, s = make([("i.xlsx", small_book(DATA), "input"), ("o.xlsx", small_book(DATA), "output")])
    with pytest.raises(ValueError, match="no sheet looks like the input"):
        service.run_system(st, s["id"], [("bad.xlsx", small_book([(1, 2, 3)], ("a", "b", "c")))])


# -------------------------------------------------------------------------------------------- toggles
def toggle_files(indicator=True):
    inp = small_book([(1, "A", 1000.0), (2, "B", 2500.0)], ("ID", "Tier", "Amount"))
    mp = openpyxl.Workbook()
    r = mp.active
    r.title = "Rates"
    r.append(["Tier", "Rate"])
    r.append(["A", 0.1])
    r.append(["B", 0.2])
    t = mp.create_sheet("Toggle")
    t.append(["Label", "Value"])
    t.append(["Indicator", indicator])
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Result"
    o.append(["ID", "Tier", "Amount", "Fee"])
    for i, (a, tier, amt) in enumerate([(1, "A", 1000.0), (2, "B", 2500.0)], start=2):
        o.append([a, tier, amt, f"=IF(Toggle!$B$2,ROUND(C{i}*VLOOKUP(B{i},Rates!$A$2:$B$3,2,FALSE),2),0)"])
    r2 = out.create_sheet("Rates")
    r2.append(["Tier", "Rate"])
    r2.append(["A", 0.1])
    r2.append(["B", 0.2])
    t2 = out.create_sheet("Toggle")
    t2.append(["Label", "Value"])
    t2.append(["Indicator", indicator])
    return inp, b(mp), b(out)


def test_toggle_and_mapping_sheets_travel_with_the_output():
    inp, mp, out = toggle_files(True)
    st, s = make([("i.xlsx", inp, "input"), ("m.xlsx", mp, "mapping"), ("o.xlsx", out, "output")])
    assert s["toggles"] == [{"name": "Indicator", "value": "Yes", "sheet": "Toggle", "cell": "B2"}]
    assert set(s["tables"]) == {"Rates", "Toggle"}
    run, wb = output_book(st, s, inp)
    assert wb["Result"]["D2"].value == "=IF(Toggle!$B$2,ROUND(C2*VLOOKUP(B2,Rates!$A$2:$B$3,2,FALSE),2),0)"
    assert wb["Toggle"]["B2"].value is True and wb["Rates"]["B3"].value == 0.2


def test_toggle_value_comes_from_the_mapping_and_can_be_overridden_per_run():
    inp, mp, out = toggle_files(True)
    st, s = make([("i.xlsx", inp, "input"), ("m.xlsx", mp, "mapping"), ("o.xlsx", out, "output")])
    run, wb = output_book(st, s, inp, mapping_files=[("m2.xlsx", toggle_files(False)[1])])
    assert wb["Toggle"]["B2"].value is False and run["toggles"][0]["value"] == "No"            # attached mapping wins
    run, wb = output_book(st, s, inp, toggle_values={"Indicator": "No"})
    assert wb["Toggle"]["B2"].value is False                                                    # per-run override
    run, wb = output_book(st, s, inp)
    assert wb["Toggle"]["B2"].value is True                                                     # stored value otherwise


def test_users_toggle_sample_and_headerless_toggle_sheet():
    f = ROOT / "dev" / "sample_data" / "toggle" / "Toggle_sample.xlsx"
    sheets = read_workbook(f.read_bytes(), f.name)
    assert sheets[0]["grid"] == [["Label", "Value"], ["Indicator", True]]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Toggle"
    ws.append(["Include Tax", "Yes"])
    ws.append(["Round Up", "No"])
    s = read_workbook(b(wb), "t.xlsx")[0]
    assert len(s["df"]) == 2 and s["grid"][0] == ["Include Tax", "Yes"]
    t = engine.toggles_from_tables({"Toggle": engine.df_to_spec(s["df"], s["header_row"], s["col_offset"], s["grid"])})
    assert [(x["name"], x["value"], x["cell"]) for x in t] == [("Include Tax", "Yes", "B1"), ("Round Up", "No", "B2")]


# ---------------------------------------------------- a sheet that cannot be lined up with the input is not fatal
def test_sheet_with_more_rows_than_the_input_gets_all_rows_and_a_warning():
    """Input has 3 rows, the sample sheet 25 (e.g. a sheet that repeats/explodes rows): must still build."""
    inp = small_book([(1, "x", 10), (2, "y", 20), (3, "x", 30)])
    out = openpyxl.Workbook()
    o = out.active
    o.title = "ABC"
    o.append(["Grp", "Amt", "Double"])
    for i in range(2, 27):                                   # 25 rows, no ID column to match them with the input
        o.append(["x", 10 * i, f"=B{i}*2"])
    st, s = make([("i.xlsx", inp, "input"), ("o.xlsx", b(out), "output")])
    assert any("could not tell which input rows" in w for w in s["warnings"])
    run, wb = output_book(st, s, small_book([(7, "x", 5), (8, "y", 6)]))
    assert run["sheets"] == {"ABC": 2} and wb["ABC"]["C2"].value == "=B2*2"


def test_one_unmatchable_sheet_does_not_stop_the_other_sheets_filters():
    rows = [(1, "x", 10), (2, "y", 20), (3, "x", 30), (4, "y", 5)]
    out = openpyxl.Workbook()
    out.remove(out.active)
    for g in ("x", "y"):
        w = out.create_sheet("Grp_" + g)
        w.append(["ID", "Grp", "Amt"])
        for r in rows:
            if r[1] == g:
                w.append(list(r))
    w = out.create_sheet("Summary")                          # 12 rows, no way to match them
    w.append(["Grp", "Amt"])
    for i in range(12):
        w.append(["x", i])
    st, s = make([("i.xlsx", small_book(rows), "input"), ("o.xlsx", b(out), "output")])
    conds = {x["name"]: x["conditions"] for x in s["filters"]["split"]["sheets"]}
    assert conds["Summary"] == [] and conds["Grp_x"] and conds["Grp_y"]            # the others still got real filters
    run, wb = output_book(st, s, small_book(rows))
    assert run["sheets"] == {"Grp_x": 2, "Grp_y": 2, "Summary": 4}


def test_sheet_with_no_input_columns_is_not_fatal():
    out = openpyxl.Workbook()
    o = out.active
    o.title = "OnlyFormulas"
    o.append(["Metric", "Value"])
    o.append(["Total rows", 4])
    o.append(["Total amount", 65])
    st, s = make([("i.xlsx", small_book(DATA), "input"), ("o.xlsx", b(out), "output")])
    assert any("none of its columns match" in w for w in s["warnings"])


def test_no_id_column_needed_to_work_out_the_sheet_filters():
    rows = [("x", 10), ("y", 20), ("x", 30), ("y", 5), ("x", 40)]
    inp = small_book(rows, ("Grp", "Amt"))
    out = openpyxl.Workbook()
    out.remove(out.active)
    for g in ("x", "y"):
        w = out.create_sheet("Only_" + g)
        w.append(["Grp", "Amt", "Tax"])
        for i, r in enumerate([r for r in rows if r[0] == g], start=2):
            w.append([r[0], r[1], f"=B{i}*0.2"])
    st, s = make([("i.xlsx", inp, "input"), ("o.xlsx", b(out), "output")])
    assert s["filters"]["split"]["column"] == "Grp" and not any("could not tell" in w for w in s["warnings"])
    run, wb = output_book(st, s, small_book([("y", 1), ("x", 2), ("x", 3)], ("Grp", "Amt")))
    assert run["sheets"] == {"Only_x": 2, "Only_y": 1} and wb["Only_x"]["C3"].value == "=B3*0.2"
