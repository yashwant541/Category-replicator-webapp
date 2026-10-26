"""Reading uploaded files into tidy DataFrames and (de)serialising mapping tables."""

from __future__ import annotations

import datetime as dt
import io
import math
import re
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd

from replicator_formulas.sample_parser import detect_header_row, _headers_from_row, _norm, is_toggle_sheet
from .rules import is_blank, nk


def jsonable(v: Any) -> Any:
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        f = float(v)
        return None if math.isnan(f) or math.isinf(f) else (int(f) if f.is_integer() and abs(f) < 1e15 else f)
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, pd.Timestamp):
        return v.isoformat()
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    return v if isinstance(v, (int, str)) else str(v)


def _rows_to_df(rows: List[List[Any]]) -> Tuple[pd.DataFrame, int]:
    rows = [[None if (isinstance(c, float) and math.isnan(c)) else c for c in r] for r in rows]
    if not any(any(not is_blank(c) for c in r) for r in rows):
        return pd.DataFrame(), 0
    hr = detect_header_row(rows)
    raw_hdr = rows[hr]
    body = rows[hr + 1:]
    width = max([len(raw_hdr)] + [len(r) for r in body[:2000]]) if body else len(raw_hdr)
    raw_hdr = list(raw_hdr) + [None] * (width - len(raw_hdr))
    keep = [j for j in range(width)
            if _norm(raw_hdr[j]) or any(j < len(r) and not is_blank(r[j]) for r in body)]
    headers = _headers_from_row([raw_hdr[j] for j in keep])
    data = [[(r[j] if j < len(r) else None) for j in keep] for r in body
            if any(not is_blank(c) for c in r)]
    df = pd.DataFrame(data, columns=headers, dtype=object)
    df.attrs["col_offset"] = int(min(keep)) if keep else 0
    return df, hr


def read_workbook(file_bytes: bytes, filename: str, report: Optional[List[Dict[str, Any]]] = None) -> List[Dict[str, Any]]:
    """-> [{sheet, header_row, df, n_formulas}] for every non-empty sheet.

    ``report`` (optional list) receives one {sheet, issue} entry per sheet that was skipped or looks wrong."""
    low = (filename or "").lower()
    out: List[Dict[str, Any]] = []
    if low.endswith((".csv", ".txt")):
        try:
            raw = pd.read_csv(io.BytesIO(file_bytes), header=None, dtype=object, keep_default_na=False,
                              skip_blank_lines=False, encoding_errors="replace")
        except TypeError:                      # pandas < 1.3 has no encoding_errors
            raw = pd.read_csv(io.BytesIO(file_bytes), header=None, dtype=object, keep_default_na=False,
                              skip_blank_lines=False, encoding="utf-8")
        df, hr = _rows_to_df([[None if (isinstance(c, str) and c == "") else c for c in r] for r in raw.values.tolist()])
        if not df.empty:
            out.append({"sheet": "Sheet1", "header_row": hr, "df": df, "n_formulas": 0})
        return out
    if low.endswith(".xls"):
        book = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, header=None,
                             dtype=object, engine="xlrd")
        for name, raw in book.items():
            df, hr = _rows_to_df(raw.values.tolist())
            if not df.empty:
                out.append({"sheet": name, "header_row": hr, "df": df, "n_formulas": 0})
        return out
    # Read once in streaming mode.  Formulas come back as "=..." strings; only if the workbook has any do
    # we read it a second time for the saved results (so plain data files are parsed just once).
    wb_f = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False, read_only=True)
    raw: Dict[str, List[List[Any]]] = {}
    for ws in wb_f.worksheets:
        try:
            ws.reset_dimensions()          # some generators write a wrong <dimension>; recompute it
        except Exception:  # noqa: BLE001
            pass
        raw[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb_f.close()

    def has_formula(rows):
        return any(isinstance(c, str) and c.startswith("=") for r in rows for c in r)
    needs_values = {n for n, rows in raw.items() if has_formula(rows)}
    vals: Dict[str, List[List[Any]]] = {}
    if needs_values:
        wb_v = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        for ws in wb_v.worksheets:
            if ws.title in needs_values:
                try:
                    ws.reset_dimensions()
                except Exception:  # noqa: BLE001
                    pass
                vals[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
        wb_v.close()
    for name, frows in raw.items():
        nf = unc = 0
        if name in needs_values:
            vrows = vals.get(name, [])
            for fr, vr in zip(frows, vrows):
                for c, v in zip(fr, vr):
                    if isinstance(c, str) and c.startswith("="):
                        nf += 1
                        unc += v is None
            rows = vrows
        else:
            rows = frows
        grid = None
        if is_toggle_sheet(name):            # name|Yes/No rows: keep the grid so cell addresses (B2) stay exact
            grid = [[jsonable(c) for c in r] for r in rows]
            while grid and not any(c is not None for c in grid[-1]):
                grid.pop()
            if grid and grid[0] and len(grid[0]) > 1 and _yn_like(grid[0][1]):     # no header row at all
                rows = [["Label", "Value"]] + [list(r[:2]) for r in rows]
        df, hr = _rows_to_df(rows)
        if df.empty:
            if report is not None:
                report.append({"sheet": name, "issue": (
                    "skipped: %d formula cell(s) have no saved result, so there are no values to read - open the "
                    "file in Excel, save it, and upload it again" % nf) if nf else "skipped: the sheet is empty"})
            continue
        if unc and report is not None:
            report.append({"sheet": name, "issue": "%d of %d formula cell(s) have no saved result and read as blank "
                           "- open the file in Excel and save it, then upload again" % (unc, nf)})
        out.append({"sheet": name, "header_row": hr, "df": df, "n_formulas": nf, "col_offset": df.attrs.get("col_offset", 0),
                    "grid": grid})
    return out


def _yn_like(v: Any) -> bool:
    return str(v).strip().lower() in ("yes", "no", "y", "n", "true", "false", "1", "0")


def yes_no(raw: Any) -> Optional[str]:
    """Toggle cell value -> 'Yes' / 'No' (None when it isn't a switch value)."""
    s = str(raw).strip().lower() if raw is not None else ""
    if s in ("yes", "y", "true", "1"):
        return "Yes"
    if s in ("no", "n", "false", "0"):
        return "No"
    return None


def guess_role(filename: str, sheet: str, n_formulas: int, is_first: bool) -> str:
    if re.search(r"dictionary|readme|read_me|notes?\b|instructions|legend|glossary|about", sheet.lower()):
        return "ignore"
    s = f"{filename} {sheet}".lower()
    if re.search(r"\b(map|mapping|lookup|rate|rates|ref|reference|master|param|toggle)", s):
        return "mapping"
    return "input"


def df_to_spec(df: pd.DataFrame, header_row: int = 0, col_offset: int = 0, grid: Any = None) -> Dict[str, Any]:
    recs = [{c: jsonable(v) for c, v in zip(df.columns, row)}
            for row in df.itertuples(index=False, name=None)]
    return {"headers": [str(c) for c in df.columns], "records": recs, "header_row": int(header_row), "col_offset": int(col_offset),
            **({"grid": grid} if grid else {})}


def spec_to_df(spec: Dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(spec["records"], columns=spec["headers"], dtype=object)


def summarize(df: pd.DataFrame, n_preview: int = 5) -> Dict[str, Any]:
    prev = [[jsonable(v) for v in row] for row in df.head(n_preview).itertuples(index=False, name=None)]
    return {"rows": int(len(df)), "cols": int(df.shape[1]),
            "headers": [str(c) for c in df.columns], "preview": prev}


def canonical_columns(df: pd.DataFrame, wanted: List[str]) -> pd.DataFrame:
    """Rename df columns to the stored names (case/space-insensitive)."""
    low = {nk(c): c for c in wanted}
    ren, seen = {}, set()
    for c in df.columns:
        t = low.get(nk(c))
        if t and t not in seen:
            ren[c] = t
            seen.add(t)
    return df.rename(columns=ren)


_SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!")


def formula_sheet_refs(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Which sheets hold formulas, and which sheets those formulas point at (mapping sheets)."""
    if not (filename or "").lower().endswith((".xlsx", ".xlsm")):
        return {"with_formulas": set(), "referenced": set()}
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    withf, refd = set(), set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    withf.add(ws.title)
                    for m in _SHEET_REF.finditer(c.value):
                        refd.add((m.group(1) or m.group(2)).replace("''", "'"))
    return {"with_formulas": withf, "referenced": refd - {ws.title for ws in wb.worksheets if ws.title in withf}}
