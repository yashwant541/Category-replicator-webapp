"""Build a *system* (a learned category) from example files, and run it on new data.

``build_system``  : input + mapping + expected output  ->  system dict (rules, tables, filters)
``apply_system``  : system + new input DataFrame      ->  output frame split into sheets
``write_workbook``: sheets -> .xlsx (values, or live Excel formulas)
``compare_frames``: produced vs expected output -> per-column match %

Dataiku-free; the backend only moves bytes in and out of managed folders.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import uuid
from collections import OrderedDict
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import rules as R
from .learn import Learner, align, build_tree, PERFECT
from .rules import Ctx, RuleError, apply_rule, cmp_mask, disp, eq_series, is_blank, nk
from .tables import canonical_columns, df_to_spec, jsonable, spec_to_df

ENGINE_BUILD = "2026.10a"
SRC = "Source sheet"
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


# =============================================================== conditions
def conditions_mask(df: pd.DataFrame, conds: List[Dict[str, Any]], match: str = "all") -> pd.Series:
    if not conds:
        return pd.Series(True, index=df.index)
    masks = []
    for c in conds:
        col = R.resolve_col(df, c["col"])
        masks.append(cmp_mask(df[col], c["op"], c.get("value")))
    m = masks[0]
    for x in masks[1:]:
        m = (m & x) if match == "all" else (m | x)
    return m


_INV = {">=": "<", "<": ">=", ">": "<=", "<=": ">", "==": "!=", "!=": "=="}


def _tree_to_conditions(tree: Dict[str, Any], keep: str = "Yes") -> Optional[List[Dict[str, Any]]]:
    out: List[Dict[str, Any]] = []
    node = tree
    while "test" in node:
        t = dict(node["test"])
        a, b = node["then"], node["else"]

        def is_leaf(n, val):
            return "leaf" in n and nk(n["leaf"].get("const")) == nk(val)
        if is_leaf(b, "No") and not is_leaf(a, "No"):
            out.append(t)
            node = a
        elif is_leaf(a, "No") and not is_leaf(b, "No"):
            t["op"] = _INV[t["op"]]
            out.append(t)
            node = b
        else:
            return None
    return out if is_leaf_yes(node) else None


def is_leaf_yes(node) -> bool:
    return "leaf" in node and nk(node["leaf"].get("const")) == "yes"


def detect_pre_filter(X: pd.DataFrame, included: np.ndarray) -> Tuple[List[Dict[str, Any]], Optional[str]]:
    if included.all():
        return [], None
    y = pd.Series(np.where(included, "Yes", "No"))
    tree = build_tree(X.reset_index(drop=True), y, max_depth=3, max_leaves=6)
    conds = _tree_to_conditions(tree) if tree else None
    if conds and conditions_mask(X.reset_index(drop=True), conds).eq(included).all():
        return conds, None
    return [], ("Some input rows are not in the output but no simple rule explains which; add an "
                "input filter manually if rows should be dropped.")


def detect_split(work: pd.DataFrame, sheet: pd.Series) -> Tuple[Dict[str, Any], Optional[str]]:
    sheets = list(dict.fromkeys(sheet))
    off = {"enabled": False, "mode": "column", "column": "", "map": {}, "new_values": "create",
           "sheets": [], "rest_sheet": ""}
    if len(sheets) < 2:
        return off, None
    for c in work.columns:
        if c == "__sheet__":
            continue
        keys = work[c].map(nk)
        modal: Dict[str, str] = {}
        ok = True
        for sh in sheets:
            vc = keys[(sheet == sh).values].value_counts()
            if vc.empty or vc.iloc[0] < 0.98 * vc.sum() or vc.index[0] == "":
                ok = False
                break
            modal[sh] = vc.index[0]
        if ok and len(set(modal.values())) == len(sheets):
            mp = {}
            for sh in sheets:
                row = int(np.argmax(((sheet == sh).values) & (keys == modal[sh]).values))
                mp[str(jsonable(work[c].iloc[row]) if not isinstance(work[c].iloc[row], str)
                       else work[c].iloc[row])] = sh
            return {"enabled": True, "mode": "column", "column": c, "map": mp,
                    "new_values": "create", "sheets": [], "rest_sheet": ""}, None
    cust = dict(off)
    cust.update(enabled=True, mode="custom",
                sheets=[{"name": sh, "match": "all", "conditions": []} for sh in sheets])
    return cust, ("The output has several sheets but no single column explains the split - define "
                  "each sheet's filter under Filters.")


def _align_sheets(X, Y, sheet_col, sheet_columns, out_group):
    """Line each output sheet up with the input on its own, so sheets may overlap (an audit sheet that
    repeats every row) or hold different subsets.  Falls back to aligning the stacked output."""
    if len(out_group) > 1:
        parts, masks = [], {}
        try:
            for e in out_group:
                rows = (sheet_col == e["sheet"]).values
                Ys = Y.loc[rows, sheet_columns[e["sheet"]]].reset_index(drop=True)
                Xs, m, inf = align(X, Ys)
                parts.append(Xs)
                masks[e["sheet"]] = m
            Xa = pd.concat(parts, ignore_index=True)
            included = np.zeros(len(X), bool)
            for m in masks.values():
                included |= m
            return Xa, included, {"method": "per sheet", "input_rows": len(X), "output_rows": len(Y)}, masks
        except ValueError:
            pass
    Xa, included, ainfo = align(X, Y)
    return Xa, included, ainfo, {}


# ================================================================ building
def detect_routing(X: pd.DataFrame, masks: Dict[str, np.ndarray], rules: List[Dict[str, Any]],
                   tables: Dict[str, pd.DataFrame], toggles: List[Dict[str, Any]]
                   ) -> Tuple[Dict[str, Any], Optional[str]]:
    """Work out, for every output sheet, which input rows it contains.

    The condition may use a *derived* column (e.g. the mapping's Output_Sheet), so the learned rules are
    run over the whole input first.  A sheet holding every row gets no condition."""
    sp = {"enabled": True, "mode": "custom", "column": "", "map": {}, "new_values": "create",
          "sheets": [], "rest_sheet": ""}
    full, _ = run_rules({"rules": [{k: v for k, v in r.items() if k != "only_in"} for r in rules], "filters": {}},
                        X.reset_index(drop=True), Ctx(tables, toggles))
    miss = []
    for name, m in masks.items():
        if m.all():
            conds: Optional[List[Dict[str, Any]]] = []
        else:
            conds = _route_conditions(full, m, name)
        if conds is None:
            miss.append(name)
            conds = []
        sp["sheets"].append({"name": name, "match": "all", "conditions": conds})
    return sp, (("Could not work out which rows belong on sheet(s) " + ", ".join(miss) +
                 " - set their filters under Filters & outputs.") if miss else None)


def _route_conditions(full: pd.DataFrame, mask: np.ndarray, sheet: str) -> Optional[List[Dict[str, Any]]]:
    best = None
    for pos, c in enumerate(full.columns):
        ks = full[c].map(nk)
        inc, exc = set(ks[mask]), set(ks[~mask])
        if not inc or "" in inc or inc & exc:
            continue
        exact = inc == {nk(sheet)}
        score = (0 if exact else 1, len(inc), pos)
        if best is None or score < best[0]:
            orig = [next(v for v, k in zip(full[c], ks) if k == i) for i in sorted(inc)]
            cond = {"col": c, "op": "==", "value": jsonable(orig[0])} if len(orig) == 1 else \
                {"col": c, "op": "in", "value": [str(jsonable(o)) for o in orig]}
            best = (score, [cond])
    if best:
        return best[1]
    tree = build_tree(full, pd.Series(np.where(mask, "Yes", "No")), max_depth=3, max_leaves=6)
    conds = _tree_to_conditions(tree) if tree else None
    return conds


def _stem(fn: str) -> str:
    return re.sub(r"\.[^.]+$", "", fn or "")


def _same_cols(a: pd.DataFrame, b: pd.DataFrame) -> bool:
    return [nk(c) for c in a.columns] == [nk(c) for c in b.columns]


def _unique_name(base: str, taken: set) -> str:
    n, i = base, 2
    while n in taken:
        n = f"{base}_{i}"
        i += 1
    return n


def build_system(entries: List[Dict[str, Any]], options: Dict[str, Any]) -> Dict[str, Any]:
    """``entries``: [{file, sheet, role, df, header_row, n_formulas, bytes}]."""
    warnings: List[str] = []
    ins = [e for e in entries if e["role"] == "input"]
    maps = [e for e in entries if e["role"] == "mapping"]
    outs = [e for e in entries if e["role"] == "output"]
    if not ins:
        raise ValueError("Add at least one input sheet.")
    if not outs:
        raise ValueError("Add the expected output sheet(s).")

    # ---- inputs: stack sheets that share a schema, the biggest stack is the primary table
    groups: List[List[Dict[str, Any]]] = []
    for e in ins:
        for g in groups:
            if _same_cols(g[0]["df"], e["df"]):
                g.append(e)
                break
        else:
            groups.append([e])
    groups.sort(key=lambda g: -sum(len(x["df"]) for x in g))
    primary = groups[0]
    parts = []
    names = [e["sheet"] for e in primary]
    for e in primary:
        f = canonical_columns(e["df"], list(primary[0]["df"].columns)).copy()
        if len(primary) > 1:   # the input's own sheets are a split: remember which sheet a row came from
            f[SRC] = e["sheet"] if names.count(e["sheet"]) == 1 else "%s/%s" % (_stem(e["file"]), e["sheet"])
        parts.append(f)
    X = pd.concat(parts, ignore_index=True)
    taken: set = set()
    tables: Dict[str, Dict[str, Any]] = {}
    table_src: Dict[str, str] = {}
    out_files = {o["file"] for o in outs}
    ordered = sorted(maps, key=lambda e: e["file"] in out_files) + [x for g in groups[1:] for x in g]
    for e in ordered:
        if e["file"] in out_files and any(nk(e["sheet"]) == nk(t) for t in tables):
            # same mapping sheet is also in the uploaded mapping file: the uploaded one wins
            mine = next(t for t in tables if nk(t) == nk(e["sheet"]))
            same = tables[mine]["headers"] == [str(c) for c in e["df"].columns] and                 len(tables[mine]["records"]) == len(e["df"])
            if not same:
                warnings.append(f"Mapping '{e['sheet']}' in the output file differs from the uploaded mapping; "
                                "the uploaded mapping is used.")
            continue
        nm = _unique_name(e["sheet"] if e["sheet"] not in taken else f"{_stem(e['file'])}_{e['sheet']}", taken)
        taken.add(nm)
        tables[nm] = df_to_spec(e["df"], e.get("header_row", 0))
        table_src[nm] = f"{e['file']} / {e['sheet']}"
    if len(groups) > 1:
        warnings.append("Input sheets with a different layout were kept as reference tables: "
                        + ", ".join(x["sheet"] for g in groups[1:] for x in g))

    # ---- output(s): every output sheet is kept; each may have its own derived columns
    out_group = list(outs)
    union: List[str] = []
    seen_nk: Dict[str, str] = {}
    sheet_columns: Dict[str, List[str]] = {}
    for e in out_group:
        cols_e = []
        for c in e["df"].columns:
            c = str(c)
            if nk(c) not in seen_nk:
                seen_nk[nk(c)] = c
                union.append(c)
            cols_e.append(seen_nk[nk(c)])
        sheet_columns[e["sheet"]] = cols_e
    frames = []
    for e in out_group:
        f = e["df"].copy()
        f.columns = sheet_columns[e["sheet"]]
        f = f.loc[:, ~f.columns.duplicated()].reindex(columns=union)
        f["__sheet__"] = e["sheet"]
        frames.append(f)
    Yall = pd.concat(frames, ignore_index=True)
    sheet_col = Yall.pop("__sheet__")
    Y = Yall
    out_cols = list(union)
    if len({tuple(v) for v in sheet_columns.values()}) > 1:
        only = [c for c in out_cols if any(c not in v for v in sheet_columns.values())]
        warnings.append("Output sheets have different columns (" + ", ".join(only[:8]) +
                        ("…" if len(only) > 8 else "") + "): each sheet keeps its own derived columns.")

    # ---- formulas in the output workbook (every data sheet is read, not only the first)
    formulas: Dict[str, Dict[str, Any]] = {}
    variants: Dict[str, Dict[str, Dict[str, Any]]] = {}
    toggles: List[Dict[str, Any]] = []
    fes = [e for e in out_group if e.get("n_formulas") and e.get("bytes")]
    if fes:
        from replicator_formulas.sample_parser import analyze_sample, sheet_to_records
        per_col: Dict[str, Dict[str, Dict[str, Any]]] = {}
        need: set = set()
        for fe in fes:
            try:
                an = analyze_sample(fe["bytes"], data_sheet=fe["sheet"])
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"Could not read formulas on sheet '{fe['sheet']}': {exc}")
                continue
            for t in an.get("toggles", []):
                if not any(x["name"] == t["name"] for x in toggles):
                    toggles.append(t)
            warnings += [w for w in an.get("warnings", []) if "No computed" not in w and w not in warnings]
            for c in an.get("computed_columns", []):
                per_col.setdefault(c["name"], {})[fe["sheet"]] = c
                need |= set(c["refs"]["sheets"])
        for col, by in per_col.items():
            if len({c["excel_formula"] for c in by.values()}) == 1:
                formulas[col] = next(iter(by.values()))
            else:
                variants[col] = by
        if variants:
            warnings.append("The formula differs between output sheets for: " + ", ".join(variants)
                            + " - each sheet's own formula is used (needs a split to tell the sheets apart).")
        have = {n.lower() for n in tables}
        src = fes[0]
        for sname in sorted(need):
            if sname.lower() not in have:
                try:
                    h, recs = sheet_to_records(src["bytes"], sname)
                    tables[sname] = {"headers": h, "records": [{k: jsonable(v) for k, v in r.items()} for r in recs],
                                     "header_row": 0}
                    table_src[sname] = f"{src['file']} / {sname} (output workbook)"
                except Exception:  # noqa: BLE001
                    warnings.append(f"Formulas reference sheet '{sname}' which was not provided as a mapping.")
    ndf = {n: spec_to_df(s) for n, s in tables.items()}

    # ---- align + learn
    Xa, included, ainfo, sheet_masks = _align_sheets(X, Y, sheet_col, sheet_columns, out_group)
    if ainfo.get("warning"):
        warnings.append(ainfo["warning"])
    y_known = bool(Y.map(lambda v: not is_blank(v)).values.any())
    if not y_known and not formulas:
        raise ValueError("The output sheet has no values (and no formulas) to learn from.")
    learner = Learner(Xa, Y, ndf, formulas, toggles, y_known, variants, sheet_col, sheet_columns)
    rule_list = learner.learn(out_cols)
    for r in rule_list:
        have = [sh for sh, cs in sheet_columns.items() if r["name"] in cs]
        if len(have) < len(sheet_columns):
            r["only_in"] = have

    # ---- filters
    flt: Dict[str, Any] = {"pre": [], "split": {"enabled": False, "mode": "column", "column": "", "map": {},
                                                "new_values": "create", "sheets": [], "rest_sheet": ""}}
    sp_opt = options.get("split") or {}
    mode = sp_opt.get("mode", "auto")
    w2 = learner.work.copy()
    pre, pre_warn = detect_pre_filter(X, included)
    flt["pre"] = pre
    if pre_warn:
        warnings.append(pre_warn)
    if mode == "none":
        pass
    elif mode == "column" and sp_opt.get("column"):
        col = sp_opt["column"]
        vals = list(dict.fromkeys(w2[col].map(lambda v: v if isinstance(v, str) else jsonable(v))))
        flt["split"].update(enabled=True, mode="column", column=col,
                            map={str(v): disp(v) for v in vals if not is_blank(v)})
    elif mode == "custom":
        flt["split"].update(enabled=True, mode="custom", sheets=sp_opt.get("sheets", []))
    else:
        sp, sw = detect_split(w2, sheet_col)
        overlap = bool(sheet_masks) and any(
            (a & b).any() for i, a in enumerate(sheet_masks.values()) for b in list(sheet_masks.values())[i + 1:])
        if len(sheet_masks) > 1 and (overlap or not sp.get("enabled") or sp.get("mode") != "column"):
            sp, sw = detect_routing(X, sheet_masks, rule_list, ndf, toggles)
        flt["split"] = sp
        if sw:
            warnings.append(sw)

    # ---- assemble
    in_cols = [str(c) for c in X.columns]
    system = {
        "engine": ENGINE_BUILD,
        "id": options.get("id") or uuid.uuid4().hex[:10],
        "name": options.get("name") or "Untitled system",
        "description": options.get("description", ""),
        "created": dt.datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "input": {"columns": in_cols, "sources": [{"file": e["file"], "sheet": e["sheet"], "rows": int(len(e["df"]))}
                                                   for e in primary]},
        "tables": tables, "table_sources": table_src,
        "output": {"columns": out_cols, "sheet_name": out_group[0]["sheet"], "sheet_columns": sheet_columns},
        "rules": rule_list, "toggles": toggles, "filters": flt,
        "training": {"align": ainfo, "output_rows": int(len(Y)), "sheets": list(dict.fromkeys(sheet_col))},
        "warnings": warnings,
    }
    finalize_system(system)
    return system


def finalize_system(system: Dict[str, Any]) -> None:
    """Recompute derived metadata (required input columns, descriptions, rule->excel)."""
    incols = {nk(c): c for c in system["input"]["columns"]}
    produced: set = set()
    required: List[str] = []
    for r in system["rules"]:
        for c in R.rule_columns(r):
            if nk(c) in incols and nk(c) not in produced and incols[nk(c)] not in required:
                required.append(incols[nk(c)])
        if not (r["kind"] == "copy" and nk(r["col"]) == nk(r["name"])):
            produced.add(nk(r["name"]))
        r["description"] = R.describe(r)
    f = system.get("filters", {})
    sp0 = f.get("split", {})
    for r in system["rules"]:
        r.pop("_pending_group", None)
        if r.get("by_sheet") and not r.get("group"):
            if sp0.get("enabled") and sp0.get("mode") == "column" and sp0.get("column"):
                r["group"] = {"col": sp0["column"], "map": {nk(k): v for k, v in sp0["map"].items()}}
            else:
                r["status"] = "approx"
                w = f"'{r['name']}': formula differs per sheet but there is no split column to tell the sheets apart."
                if w not in system.setdefault("warnings", []):
                    system["warnings"].append(w)
        if r.get("group") and nk(r["group"]["col"]) in incols and incols[nk(r["group"]["col"])] not in required:
            required.append(incols[nk(r["group"]["col"])])
    for c in f.get("pre", []):
        if nk(c["col"]) in incols and incols[nk(c["col"])] not in required:
            required.append(incols[nk(c["col"])])
    sp = f.get("split", {})
    cols = [sp.get("column")] if sp.get("enabled") and sp.get("mode") == "column" else []
    for s in sp.get("sheets", []) if sp.get("enabled") else []:
        cols += [c["col"] for c in s.get("conditions", [])]
    for c in cols:
        if c and nk(c) in incols and incols[nk(c)] not in required:
            required.append(incols[nk(c)])
    system["input"]["required"] = required


# ================================================================ applying
def tables_to_frames(system: Dict[str, Any], override: Optional[Dict[str, pd.DataFrame]] = None
                     ) -> Dict[str, pd.DataFrame]:
    t = {n: spec_to_df(s) for n, s in system.get("tables", {}).items()}
    for n, df in (override or {}).items():
        t[n] = df
    return t


def prepare_input(system: Dict[str, Any], df: pd.DataFrame) -> Tuple[pd.DataFrame, List[str]]:
    wanted = system["input"]["columns"]
    df = canonical_columns(df.reset_index(drop=True), wanted)
    notes = []
    if SRC in wanted and SRC not in df.columns:
        df[SRC] = "Input"
        notes.append("The input had a single sheet, so no split by input sheet is possible (sheet name 'Input' used).")
    have = set(df.columns)
    missing_req = [c for c in system["input"].get("required", []) if c not in have]
    if missing_req:
        raise ValueError("Input is missing required column(s): " + ", ".join(missing_req) +
                         ". Expected columns: " + ", ".join(wanted))
    for c in wanted:
        if c not in have:
            df[c] = np.nan
            notes.append(f"Input column '{c}' not found - treated as blank.")
    extra = [c for c in df.columns if c not in set(wanted)]
    if extra:
        notes.append("Ignored extra input column(s): " + ", ".join(map(str, extra[:8])))
    return df, notes


def _sheet_of_rows(system: Dict[str, Any], work: pd.DataFrame) -> Optional[pd.Series]:
    """Output-sheet name each row will land in (only for a column-based split)."""
    sp = system.get("filters", {}).get("split", {})
    if not (sp.get("enabled") and sp.get("mode") == "column" and sp.get("column")):
        return None
    try:
        col = R.resolve_col(work, sp["column"])
    except R.RuleError:
        return None
    mapping = {nk(k): v for k, v in sp.get("map", {}).items()}
    return work[col].map(lambda v: mapping.get(nk(v)))


def _only_in_mask(system: Dict[str, Any], rule: Dict[str, Any], work: pd.DataFrame) -> Optional[np.ndarray]:
    if not rule.get("only_in"):
        return None
    if system.get("filters", {}).get("split", {}).get("mode") == "custom":
        return None            # sheet membership comes from conditions that may use these very columns
    grp = _sheet_of_rows(system, work)
    if grp is None:
        return None
    m = (grp.isin(rule["only_in"]) | grp.isna()).values   # unknown sheets: compute, selection drops later
    return m if not m.all() else None


def run_rules(system: Dict[str, Any], df: pd.DataFrame, ctx: Ctx,
              only_missing_ok: bool = True) -> Tuple[pd.DataFrame, Dict[str, str]]:
    work = df.copy()
    errors: Dict[str, str] = {}
    pending = list(system["rules"])
    for _ in range(4):
        progressed, still = False, []
        for r in pending:
            if r.get("disabled"):
                work[r["name"]] = np.nan
                progressed = True
                continue
            try:
                mask = _only_in_mask(system, r, work)
                if mask is None:
                    work[r["name"]] = apply_rule(r, work, ctx).values
                else:
                    col = pd.Series(np.nan, index=work.index, dtype=object)
                    if mask.any():
                        col[mask] = apply_rule(r, work[mask], ctx).values
                    work[r["name"]] = col.values
                errors.pop(r["name"], None)
                progressed = True
            except RuleError as exc:
                errors[r["name"]] = str(exc)
                still.append(r)
            except Exception as exc:  # noqa: BLE001
                errors[r["name"]] = f"{type(exc).__name__}: {exc}"
                work[r["name"]] = np.nan
                progressed = True
        pending = still
        if not pending or not progressed:
            break
    for r in pending:
        work[r["name"]] = np.nan
    return work, errors


def safe_sheet_name(name: str, taken: set) -> str:
    base = _BAD_SHEET.sub("-", str(name)).strip().strip("'") or "Sheet"
    base = base[:31]
    n, i = base, 2
    while n.lower() in {t.lower() for t in taken}:
        suf = f" ({i})"
        n = base[:31 - len(suf)] + suf
        i += 1
    taken.add(n)
    return n


def split_frames(system: Dict[str, Any], out: pd.DataFrame, work: pd.DataFrame
                 ) -> Tuple["OrderedDict[str, pd.DataFrame]", List[str]]:
    sp = system.get("filters", {}).get("split", {})
    notes: List[str] = []
    taken: set = set()
    res0: "OrderedDict[str, pd.DataFrame]" = OrderedDict()
    res = res0
    if not sp.get("enabled"):
        res[safe_sheet_name(system["output"].get("sheet_name") or "Output", taken)] = out
        return res, notes
    sc = system["output"].get("sheet_columns") or {}
    common = [c for c in system["output"]["columns"] if all(c in v for v in sc.values())] if sc else []

    def pick(name: str, df: pd.DataFrame) -> pd.DataFrame:
        """Each sheet shows its own columns; sheets the system never saw get the shared ones."""
        cols = sc.get(name) or (common if sc and common else None)
        return df[[c for c in cols if c in df.columns]] if cols else df
    if sp.get("mode") == "column":
        col = R.resolve_col(work, sp["column"])
        keys = work[col].map(nk)
        mapping = {nk(k): v for k, v in sp.get("map", {}).items()}
        sheet_of = []
        unassigned = 0
        for v, k in zip(work[col], keys):
            if k == "":
                sheet_of.append(None)
                unassigned += 1
            elif k in mapping:
                sheet_of.append(mapping[k])
            elif sp.get("new_values", "create") == "create":
                sheet_of.append(disp(v))
            else:
                sheet_of.append(None)
                unassigned += 1
        s = pd.Series(sheet_of, index=out.index, dtype=object)
        order = [v for v in dict.fromkeys(mapping.values()) if (s == v).any()] + \
                [v for v in dict.fromkeys(x for x in s if x is not None) if v not in set(mapping.values())]
        for nm in order:
            res[safe_sheet_name(nm, taken)] = out[(s == nm).values]
        if unassigned:
            notes.append(f"{unassigned} row(s) had a blank/unknown '{sp['column']}' and were not placed in a split sheet"
                         + (" (put in 'Unassigned')." if sp.get("rest_sheet") else "."))
            if sp.get("rest_sheet"):
                res[safe_sheet_name(sp["rest_sheet"] or "Unassigned", taken)] = out[s.isna().values]
        if not res:
            res[safe_sheet_name("Output", taken)] = out
        return OrderedDict((k, pick(k, v)) for k, v in res.items()), notes
    covered = pd.Series(False, index=out.index)
    for sh in sp.get("sheets", []):
        m = conditions_mask(work, sh.get("conditions", []), sh.get("match", "all"))
        covered |= m.values
        res[safe_sheet_name(sh["name"], taken)] = out[m.values]
    if sp.get("rest_sheet"):
        res[safe_sheet_name(sp["rest_sheet"], taken)] = out[~covered.values]
    elif (~covered).any():
        notes.append(f"{int((~covered).sum())} row(s) matched no sheet filter and were left out.")
    if not res:
        res[safe_sheet_name("Output", taken)] = out
    return OrderedDict((k, pick(k, v)) for k, v in res.items()), notes


def apply_system(system: Dict[str, Any], df_in: pd.DataFrame,
                 table_override: Optional[Dict[str, pd.DataFrame]] = None) -> Dict[str, Any]:
    df, notes = prepare_input(system, df_in)
    n_in = len(df)
    pre = system.get("filters", {}).get("pre", [])
    if pre:
        keep = conditions_mask(df, pre).values
        df = df[keep].reset_index(drop=True)
        notes.append(f"Input filter kept {len(df)} of {n_in} rows.")
    ctx = Ctx(tables_to_frames(system, table_override), system.get("toggles", []))
    work, errors = run_rules(system, df, ctx)
    cols = system["output"]["columns"]
    out = work[cols].copy() if len(work) else pd.DataFrame(columns=cols)
    sheets, snotes = split_frames(system, out, work)
    notes += snotes
    for nm, st in ctx.stats.items():
        notes.append(f"'{nm}': {st['unmatched']} row(s) had no match in the mapping "
                     f"(e.g. {', '.join(map(str, st['examples'][:5]))}).")
    return {"out": out, "work": work, "sheets": sheets, "notes": notes, "rule_errors": errors,
            "stats": ctx.stats, "input_rows": n_in, "rows": len(out)}


# =============================================================== comparison
def compare_frames(pred: pd.DataFrame, expected: pd.DataFrame, max_examples: int = 8) -> Dict[str, Any]:
    cols = [c for c in expected.columns if c in pred.columns]
    res: Dict[str, Any] = {"columns": [], "rows_pred": int(len(pred)), "rows_expected": int(len(expected)),
                           "aligned": True, "note": ""}
    if not len(expected) or not len(pred):
        res["aligned"] = False
        res["note"] = "Nothing to compare (empty output)."
        return res
    try:
        P, _, info = align(pred[cols].reset_index(drop=True), expected[cols].reset_index(drop=True))
        res["aligned_by"] = info.get("method")
    except ValueError as exc:
        res["aligned"] = False
        res["note"] = str(exc)
        return res
    E = expected[cols].reset_index(drop=True)
    total_ok = 0
    for c in cols:
        eq = eq_series(P[c], E[c])
        bad = np.where(~eq.values)[0]
        res["columns"].append({
            "column": c, "match": round(float(eq.mean()), 4), "mismatches": int(len(bad)),
            "examples": [{"row": int(i) + 2, "expected": jsonable(E[c].iloc[i]),
                          "got": jsonable(P[c].iloc[i])} for i in bad[:max_examples]],
        })
        total_ok += int(eq.sum())
    res["cells"] = int(len(E) * len(cols))
    res["match"] = round(total_ok / max(1, res["cells"]), 4)
    res["missing_columns"] = [c for c in expected.columns if c not in pred.columns]
    return res


def compare_sheets(pred_sheets: Dict[str, pd.DataFrame], Y_sheets: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    """Per-sheet comparison (each sheet only has its own columns; sheets may overlap)."""
    names = {nk(k): k for k in pred_sheets}
    per, cells, ok, cols_acc, issues = {}, 0, 0, {}, []
    for k, exp in Y_sheets.items():
        if nk(k) not in names:
            issues.append(f"{k}: sheet was not produced")
            continue
        c = compare_frames(pred_sheets[names[nk(k)]], exp)
        per[k] = c
        if not c.get("aligned"):
            issues.append(f"{k}: {c.get('note')}")
            continue
        cells += c["cells"]
        ok += round(c["match"] * c["cells"])
        for col in c["columns"]:
            a = cols_acc.setdefault(col["column"], {"column": col["column"], "n": 0, "bad": 0, "examples": []})
            a["n"] += len(exp)
            a["bad"] += col["mismatches"]
            a["examples"] += [dict(e, sheet=k) for e in col["examples"]][:4]
    extra = [k for k in pred_sheets if nk(k) not in {nk(x) for x in Y_sheets}]
    return {"aligned": bool(cells), "note": "; ".join(issues), "cells": cells,
            "match": round(ok / max(1, cells), 4),
            "rows_pred": int(sum(len(v) for v in pred_sheets.values())),
            "rows_expected": int(sum(len(v) for v in Y_sheets.values())),
            "extra_sheets": extra,
            "per_sheet": {k: {"match": v.get("match"), "rows": v.get("rows_pred")} for k, v in per.items()},
            "columns": [{"column": a["column"], "match": round(1 - a["bad"] / max(1, a["n"]), 4),
                         "mismatches": a["bad"], "examples": a["examples"][:8]} for a in cols_acc.values()]}


def replay_training(system: Dict[str, Any], X: pd.DataFrame, Y: pd.DataFrame,
                    Y_sheets: Optional[Dict[str, pd.DataFrame]] = None) -> Dict[str, Any]:
    """Run the system on the training input and diff it with the expected output (per sheet)."""
    r = apply_system(system, X)
    names = {nk(k) for k in r["sheets"]}
    if Y_sheets and all(nk(k) in names for k in Y_sheets):
        res = compare_sheets(r["sheets"], Y_sheets)
    else:
        res = compare_frames(r["out"], Y)
    res["rule_errors"] = r["rule_errors"]
    return res


# ================================================================ workbook out
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")
HELP_FILL = PatternFill("solid", fgColor="E8EDF3")


def xl_val(v: Any) -> Any:
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, float) and np.isnan(v):
        return None
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, str) and v.startswith("="):
        return "'" + v
    return v


def _autosize(ws, header_len: List[int]) -> None:
    for i, w in enumerate(header_len, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(48, w + 2))


def write_workbook(system: Dict[str, Any], sheets: "OrderedDict[str, pd.DataFrame]",
                   work: Optional[pd.DataFrame] = None, summary: Optional[List[Tuple[str, Any]]] = None,
                   live_formulas: bool = False, tables: Optional[Dict[str, pd.DataFrame]] = None) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rules = {r["name"]: r for r in system["rules"]}
    tables = tables if tables is not None else tables_to_frames(system)
    table_sheet: Dict[str, str] = {}
    taken = set(sheets.keys()) | {"Run summary"}
    any_live = False
    incols = system["input"]["columns"]

    if live_formulas:
        used_tables: List[str] = []
        for r in system["rules"]:
            used_tables += R.rule_tables(r)
        for t in dict.fromkeys(used_tables):
            if t in tables:
                table_sheet[t] = safe_sheet_name(t, taken)

    def table_ref(t, col):
        T = tables[t]
        j = list(T.columns).index(R.resolve_col(T, col)) + 1
        L = get_column_letter(j)
        return f"{R.sheet_token(table_sheet[t])}!${L}:${L}"

    def live_for(cols: List[str]):
        """Excel formulas for this sheet's own column layout (+ helper columns for raw inputs it needs)."""
        letters = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
        helpers = []
        for r in system["rules"]:
            if r["name"] not in cols:
                continue
            for c in R.rule_columns(r):
                if c not in letters and c in incols and c not in helpers:
                    helpers.append(c)
        for i, c in enumerate(helpers):
            letters[c] = get_column_letter(len(cols) + i + 1)

        def ref(c):
            if c not in letters:
                raise KeyError(c)
            return f"{letters[c]}{{r}}"
        live: Dict[str, str] = {}
        for name in cols:
            r = rules.get(name)
            if not r or r.get("disabled") or r["kind"] == "formula" and "{tok:" in (r.get("excel_formula") or ""):
                continue
            try:
                f = R.excel_formula(r, ref, table_ref)
            except Exception:  # noqa: BLE001
                f = None
            if f:
                live[name] = f
        return live, helpers

    for sname, df in sheets.items():
        ws = wb.create_sheet(sname)
        cols = [c for c in df.columns]
        live, helpers = live_for(cols) if live_formulas else ({}, [])
        any_live = any_live or bool(live)
        headers = list(cols) + helpers
        ws.append(headers)
        for j, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=j)
            c.font = HDR_FONT
            c.fill = HDR_FILL if j <= len(cols) else PatternFill("solid", fgColor="6B7C93")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [len(str(h)) for h in headers]
        hv = None
        if helpers:
            hv = work.loc[df.index, helpers] if work is not None else None
        for ri, (idx, row) in enumerate(df.iterrows(), start=2):
            line = []
            for j, c in enumerate(cols):
                if c in live:
                    line.append(live[c].replace("{r}", str(ri)))
                else:
                    line.append(xl_val(row[c]))
            if helpers:
                line += [xl_val(hv.loc[idx, h]) if hv is not None else None for h in helpers]
            ws.append(line)
            for j, v in enumerate(line):
                if v is not None and not (isinstance(v, str) and v.startswith("=")) and ri < 200:
                    widths[j] = max(widths[j], len(str(v)[:60]))
        ws.freeze_panes = "A2"
        _autosize(ws, widths)

    for t, sn in table_sheet.items():
        ws = wb.create_sheet(sn)
        T = tables[t]
        ws.append([str(c) for c in T.columns])
        for row in T.itertuples(index=False, name=None):
            ws.append([xl_val(v) for v in row])
        for j in range(1, T.shape[1] + 1):
            ws.cell(row=1, column=j).font = Font(bold=True)

    sm = wb.create_sheet("Run summary")
    sm.append(["System", system["name"]])
    sm.append(["Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")])
    sm.append(["Formulas", "live Excel formulas (beta)" if any_live else "values"])
    sm.append([])
    sm.append(["Sheet", "Rows"])
    for c in sm[5]:
        c.font = Font(bold=True)
    for sname, df in sheets.items():
        sm.append([sname, int(len(df))])
    if summary:
        sm.append([])
        for k, v in summary:
            sm.append([k, v if not isinstance(v, (list, dict)) else "; ".join(map(str, v))])
    sm.column_dimensions["A"].width = 30
    sm.column_dimensions["B"].width = 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
